# maint_mail.py — 정비관리 설정: 협력사 접수(D-1 승인) 정비의뢰 일일 메일
from __future__ import annotations

import asyncio
import base64
import hashlib
import hmac
import json
import os
import re
import ssl
from datetime import datetime, timedelta, timezone
from email.header import Header
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.utils import formataddr
from io import BytesIO
from zoneinfo import ZoneInfo

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from models import (
    AppSetting,
    Building,
    Equipment,
    Floor,
    Partner,
    Site,
    WorkOrder,
    WorkOrderStatus,
    Zone,
)

KST = ZoneInfo("Asia/Seoul")
JOB_ID = "maint_partner_receipt_mail"
PREFIX = "maint_mail."
MAX_LOOKBACK = timedelta(days=7)

DEFAULTS: dict[str, str] = {
    "enabled": "0",
    "hour": "17",
    "minute": "0",
    "to": "",
    "per_partner": "0",
    "partner_to": "{}",
    "send_empty": "0",
    "last_sent_utc": "",
    "last_run_date": "",
    "last_result": "",
    "last_result_at": "",
}

# 공식 발신 계정 — server: 서버 환경변수(RESEND_API_KEY/SMTP_*) 사용
SENDER_DEFAULTS: dict[str, str] = {
    "sender_mode": "server",
    "sender_name": "Smart FMS",
    "sender_addr": "",
    "reply_to": "",
    "smtp_host": "smtp.gmail.com",
    "smtp_port": "465",
    "smtp_security": "ssl",
    "smtp_user": "",
    "smtp_password": "",
    "resend_key": "",
}
SENDER_MODES = ("server", "smtp", "resend")
SECRET_KEYS = ("smtp_password", "resend_key")
SMTP_PRESETS = (
    ("gmail", "Gmail", "smtp.gmail.com", 465, "ssl"),
    ("naver", "네이버", "smtp.naver.com", 465, "ssl"),
    ("daum", "다음", "smtp.daum.net", 465, "ssl"),
    ("office365", "Outlook/Office 365", "smtp.office365.com", 587, "starttls"),
)

STATUS_LABELS = {
    "received": "정비의뢰",
    "assigned": "정비의뢰",
    "in_progress": "정비진행",
    "completed": "정비완료",
    "verified": "정비완료 최종승인 대기",
    "closed": "정비완료",
}
PRIORITY_LABELS = {"urgent": "긴급", "high": "높음", "normal": "보통", "low": "낮음"}

_EMAIL_RE = re.compile(r"^[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+$")


def _truthy(v: str | None) -> bool:
    return (v or "").strip().lower() in ("1", "true", "yes", "on")


def parse_emails(raw: str | None) -> tuple[list[str], list[str]]:
    """콤마·세미콜론·줄바꿈 구분 → (유효 주소, 잘못된 항목)."""
    ok: list[str] = []
    bad: list[str] = []
    for part in re.split(r"[,;\s]+", raw or ""):
        addr = part.strip()
        if not addr:
            continue
        if _EMAIL_RE.match(addr):
            if addr.lower() not in {a.lower() for a in ok}:
                ok.append(addr)
        else:
            bad.append(addr)
    return ok, bad


async def _get(session: AsyncSession, key: str) -> str:
    row = (
        await session.execute(select(AppSetting).where(AppSetting.key == PREFIX + key))
    ).scalar_one_or_none()
    if row and row.value is not None:
        return row.value
    return DEFAULTS.get(key, SENDER_DEFAULTS.get(key, ""))


async def _set(session: AsyncSession, key: str, value: str) -> None:
    row = (
        await session.execute(select(AppSetting).where(AppSetting.key == PREFIX + key))
    ).scalar_one_or_none()
    if row:
        row.value = value
    else:
        session.add(AppSetting(key=PREFIX + key, value=value))


def _clamp_int(raw: str | None, lo: int, hi: int, default: int) -> int:
    try:
        v = int(str(raw).strip())
    except (TypeError, ValueError):
        return default
    return max(lo, min(hi, v))


async def load_settings(session: AsyncSession) -> dict:
    raw = {k: await _get(session, k) for k in DEFAULTS}
    try:
        partner_to = json.loads(raw["partner_to"] or "{}")
        if not isinstance(partner_to, dict):
            partner_to = {}
    except ValueError:
        partner_to = {}
    return {
        "enabled": _truthy(raw["enabled"]),
        "hour": _clamp_int(raw["hour"], 0, 23, 17),
        "minute": _clamp_int(raw["minute"], 0, 59, 0),
        "to": raw["to"],
        "per_partner": _truthy(raw["per_partner"]),
        "partner_to": {str(k): str(v) for k, v in partner_to.items()},
        "send_empty": _truthy(raw["send_empty"]),
        "last_sent_utc": raw["last_sent_utc"],
        "last_run_date": raw["last_run_date"],
        "last_result": raw["last_result"],
        "last_result_at": raw["last_result_at"],
    }


async def save_settings(
    session: AsyncSession,
    *,
    enabled: bool,
    hour: int,
    minute: int,
    to: str,
    per_partner: bool,
    partner_to: dict[str, str],
    send_empty: bool,
) -> None:
    await _set(session, "enabled", "1" if enabled else "0")
    await _set(session, "hour", str(max(0, min(23, hour))))
    await _set(session, "minute", str(max(0, min(59, minute))))
    await _set(session, "to", to)
    await _set(session, "per_partner", "1" if per_partner else "0")
    await _set(
        session,
        "partner_to",
        json.dumps({k: v for k, v in partner_to.items() if v.strip()}, ensure_ascii=False),
    )
    await _set(session, "send_empty", "1" if send_empty else "0")


def _parse_utc(raw: str | None) -> datetime | None:
    try:
        dt = datetime.fromisoformat((raw or "").strip())
    except ValueError:
        return None
    if dt.tzinfo is not None:
        dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
    return dt


def window_for(settings: dict, now_utc: datetime) -> tuple[datetime, datetime]:
    """직전 발송 이후 ~ 지금 (최초·장기 미발송 시 24시간 / 최대 7일)."""
    since = _parse_utc(settings.get("last_sent_utc"))
    if since is None or since >= now_utc:
        since = now_utc - timedelta(hours=24)
    since = max(since, now_utc - MAX_LOOKBACK)
    return since, now_utc


def fmt_kst(dt: datetime | None, fmt: str = "%Y-%m-%d %H:%M") -> str:
    if not isinstance(dt, datetime):
        return ""
    d = dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    return d.astimezone(KST).strftime(fmt)


async def fetch_receipts(
    session: AsyncSession, since_utc: datetime, until_utc: datetime
) -> list[dict]:
    """협력사 지정 + D-1 승인된 정비의뢰 중 승인시각이 구간 안인 건."""
    rows = (
        await session.execute(
            select(WorkOrder, Partner, Equipment, Building.name, Site.name)
            .join(Partner, WorkOrder.partner_id == Partner.id)
            .outerjoin(Equipment, WorkOrder.equipment_id == Equipment.id)
            .outerjoin(Zone, Equipment.zone_id == Zone.id)
            .outerjoin(Floor, Zone.floor_id == Floor.id)
            .outerjoin(Building, Floor.building_id == Building.id)
            .outerjoin(Site, Site.id == WorkOrder.site_id)
            .where(
                WorkOrder.is_active == True,  # noqa: E712
                WorkOrder.d1_approved == True,  # noqa: E712
                WorkOrder.approved_at.is_not(None),
                WorkOrder.approved_at >= since_utc,
                WorkOrder.approved_at < until_utc,
            )
            .order_by(Partner.name, WorkOrder.approved_at, WorkOrder.id)
        )
    ).all()
    out: list[dict] = []
    for wo, partner, eq, bname, sname in rows:
        status = wo.status.value if isinstance(wo.status, WorkOrderStatus) else str(wo.status or "")
        out.append(
            {
                "id": wo.id,
                "partner_id": partner.id,
                "partner_name": partner.name,
                "partner_email": (partner.email or "").strip(),
                "approved_at": fmt_kst(wo.approved_at),
                "approved_by": wo.approved_by or "",
                "scheduled_date": wo.scheduled_date.isoformat() if wo.scheduled_date else "",
                "priority": PRIORITY_LABELS.get(wo.priority or "", wo.priority or ""),
                "location": " / ".join(x for x in (sname, bname) if x),
                "eq_code": eq.code if eq else "",
                "eq_name": eq.name if eq else "",
                "title": (wo.title or "").strip(),
                "description": (wo.description or "").strip(),
                "work_type": wo.work_type or "",
                "requester": wo.requester_name or "",
                "status": STATUS_LABELS.get(status, status),
            }
        )
    return out


XLSX_COLUMNS = (
    ("No", "id", 8),
    ("협력사", "partner_name", 16),
    ("접수(D-1 승인)일시", "approved_at", 18),
    ("승인자", "approved_by", 14),
    ("예정일", "scheduled_date", 12),
    ("우선순위", "priority", 9),
    ("위치", "location", 22),
    ("설비코드", "eq_code", 16),
    ("설비명", "eq_name", 22),
    ("정비제목", "title", 30),
    ("정비내용", "description", 40),
    ("작업유형", "work_type", 12),
    ("의뢰자", "requester", 12),
    ("진행상태", "status", 12),
)


def build_xlsx(rows: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "협력사 접수"
    ws.append([c[0] for c in XLSX_COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0F2A5C")
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append([r.get(c[1], "") for c in XLSX_COLUMNS])
        for cell in ws[ws.max_row]:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for idx, (_, _, width) in enumerate(XLSX_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_body(rows: list[dict], since_utc: datetime, until_utc: datetime, partner_name: str = "") -> str:
    head = f"[{partner_name}] " if partner_name else ""
    lines = [
        f"{head}협력사 접수 정비의뢰 현황",
        f"집계 구간(KST): {fmt_kst(since_utc)} ~ {fmt_kst(until_utc)}",
        f"접수 건수: {len(rows)}건",
        "",
    ]
    if not rows:
        lines.append("해당 구간에 협력사로 접수된 정비의뢰가 없습니다.")
    else:
        groups: dict[str, list[dict]] = {}
        for r in rows:
            groups.setdefault(r["partner_name"], []).append(r)
        for name, items in groups.items():
            if not partner_name:
                lines.append(f"■ {name} ({len(items)}건)")
            for r in items:
                eq = " ".join(x for x in (r["eq_code"], r["eq_name"]) if x)
                meta = " · ".join(
                    x
                    for x in (
                        r["approved_at"],
                        f"예정 {r['scheduled_date']}" if r["scheduled_date"] else "",
                        r["priority"],
                        r["location"],
                        eq,
                    )
                    if x
                )
                lines.append(f"  - #{r['id']} {r['title'] or r['description'][:60]}")
                if meta:
                    lines.append(f"      {meta}")
            lines.append("")
        lines.append("상세 내용은 첨부 엑셀을 확인하세요.")
    lines += ["", "※ Smart FMS 정비관리 > 설정에서 자동 발송된 메일입니다."]
    return "\n".join(lines)


def _secret_keys() -> tuple[bytes, bytes]:
    base = (os.environ.get("APP_SECRET_KEY") or "change_this_secret_in_prod").encode()
    return (
        hmac.new(base, b"maint-mail-enc", hashlib.sha256).digest(),
        hmac.new(base, b"maint-mail-mac", hashlib.sha256).digest(),
    )


def _keystream(key: bytes, nonce: bytes, n: int) -> bytes:
    out = bytearray()
    i = 0
    while len(out) < n:
        out += hmac.new(key, nonce + i.to_bytes(4, "big"), hashlib.sha256).digest()
        i += 1
    return bytes(out[:n])


def seal_secret(plain: str) -> str:
    """비밀번호·API 키 DB 저장용 암호화 (APP_SECRET_KEY 기반, 변조 검증 포함)."""
    if not plain:
        return ""
    enc, mac = _secret_keys()
    nonce = os.urandom(16)
    data = plain.encode("utf-8")
    ct = bytes(a ^ b for a, b in zip(data, _keystream(enc, nonce, len(data))))
    tag = hmac.new(mac, nonce + ct, hashlib.sha256).digest()[:16]
    return "v1:" + base64.urlsafe_b64encode(nonce + tag + ct).decode("ascii")


def open_secret(token: str) -> str | None:
    """복호화 실패(APP_SECRET_KEY 변경 등) 시 None."""
    if not token:
        return ""
    if not token.startswith("v1:"):
        return None
    try:
        raw = base64.urlsafe_b64decode(token[3:].encode("ascii"))
    except ValueError:
        return None
    nonce, tag, ct = raw[:16], raw[16:32], raw[32:]
    enc, mac = _secret_keys()
    if not hmac.compare_digest(tag, hmac.new(mac, nonce + ct, hashlib.sha256).digest()[:16]):
        return None
    try:
        return bytes(a ^ b for a, b in zip(ct, _keystream(enc, nonce, len(ct)))).decode("utf-8")
    except UnicodeDecodeError:
        return None


async def load_sender(session: AsyncSession) -> dict:
    raw = {k: await _get(session, k) for k in SENDER_DEFAULTS}
    mode = raw["sender_mode"] if raw["sender_mode"] in SENDER_MODES else "server"
    out = {
        "mode": mode,
        "name": raw["sender_name"].strip(),
        "addr": raw["sender_addr"].strip(),
        "reply_to": raw["reply_to"].strip(),
        "smtp_host": raw["smtp_host"].strip(),
        "smtp_port": _clamp_int(raw["smtp_port"], 1, 65535, 465),
        "smtp_security": "starttls" if raw["smtp_security"] == "starttls" else "ssl",
        "smtp_user": raw["smtp_user"].strip(),
    }
    for key in SECRET_KEYS:
        val = open_secret(raw[key])
        out[key] = val or ""
        out[key + "_set"] = bool(raw[key])
        out[key + "_broken"] = val is None
    return out


def sender_status(sender: dict) -> tuple[bool, str]:
    """(발송 가능 여부, 화면 안내 문구)."""
    mode = sender["mode"]
    if mode == "smtp":
        missing = [
            label
            for label, ok in (
                ("SMTP 서버", sender["smtp_host"]),
                ("계정(아이디)", sender["smtp_user"]),
                ("비밀번호", sender["smtp_password"]),
            )
            if not ok
        ]
        if sender["smtp_password_broken"]:
            return False, "저장된 비밀번호를 읽을 수 없습니다. 비밀번호를 다시 입력하세요."
        if missing:
            return False, "공식 계정 설정이 덜 됐습니다: " + ", ".join(missing)
        return True, f"공식 계정 {sender_from(sender)} (SMTP {sender['smtp_host']})"
    if mode == "resend":
        if sender["resend_key_broken"]:
            return False, "저장된 Resend API 키를 읽을 수 없습니다. 다시 입력하세요."
        if not sender["resend_key"] or not sender["addr"]:
            return False, "Resend API 키와 발신 주소를 모두 입력하세요."
        return True, f"공식 계정 {sender_from(sender)} (Resend)"
    if os.environ.get("RESEND_API_KEY", "").strip() or os.environ.get("SMTP_HOST", "").strip():
        return True, "서버 기본 발신 설정(환경변수) 사용"
    return False, "발신 계정이 없습니다. 아래 「공식 발신 계정」을 설정하세요."


def sender_from(sender: dict) -> str:
    addr = sender["addr"] or (sender["smtp_user"] if sender["mode"] == "smtp" else "")
    name = sender["name"]
    return f"{name} <{addr}>" if name and addr else addr


def mail_configured(sender: dict) -> bool:
    return sender_status(sender)[0]


def _mime(sender: dict, to_addr: str, subject: str, body: str, data: bytes, fname: str) -> tuple[str, MIMEMultipart]:
    from_addr = sender["addr"] or sender["smtp_user"]
    msg = MIMEMultipart()
    msg["Subject"] = str(Header(subject, "utf-8"))
    msg["From"] = formataddr((sender["name"], from_addr), charset="utf-8") if sender["name"] else from_addr
    msg["To"] = to_addr
    if sender["reply_to"]:
        msg["Reply-To"] = sender["reply_to"]
    msg.attach(MIMEText(body, "plain", "utf-8"))
    part = MIMEApplication(data, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    part.add_header("Content-Disposition", "attachment", filename=("utf-8", "", fname))
    msg.attach(part)
    return from_addr, msg


def _send_smtp(sender: dict, to_addr: str, subject: str, body: str, data: bytes, fname: str) -> None:
    from streetlamp.reporting import _SMTPForceIPv4, _SMTPSSLForceIPv4

    from_addr, msg = _mime(sender, to_addr, subject, body, data, fname)
    host, port = sender["smtp_host"], sender["smtp_port"]
    timeout = 60.0
    if sender["smtp_security"] == "ssl":
        smtp = _SMTPSSLForceIPv4(host, port, timeout=timeout, context=ssl.create_default_context())
    else:
        smtp = _SMTPForceIPv4(host, port, timeout=timeout)
    with smtp:
        if sender["smtp_security"] == "starttls":
            smtp.starttls(context=ssl.create_default_context())
        smtp.login(sender["smtp_user"], "".join(sender["smtp_password"].split()))
        refused = smtp.sendmail(from_addr, [to_addr], msg.as_string())
        if refused:
            raise RuntimeError(f"수신 거부: {refused}")


def _send_resend(sender: dict, to_addr: str, subject: str, body: str, data: bytes, fname: str) -> None:
    import httpx

    payload = {
        "from": sender_from(sender),
        "to": [to_addr],
        "subject": subject,
        "text": body,
        "attachments": [{"filename": fname, "content": base64.b64encode(data).decode("ascii")}],
    }
    if sender["reply_to"]:
        payload["reply_to"] = sender["reply_to"]
    with httpx.Client(timeout=60.0) as client:
        r = client.post(
            "https://api.resend.com/emails",
            headers={"Authorization": f"Bearer {sender['resend_key']}"},
            json=payload,
        )
    if r.status_code >= 400:
        raise RuntimeError(f"Resend HTTP {r.status_code}: {r.text[:300]}")


def _send_sync(sender: dict, to_addr: str, subject: str, body: str, data: bytes, fname: str) -> None:
    if sender["mode"] == "smtp":
        _send_smtp(sender, to_addr, subject, body, data, fname)
    elif sender["mode"] == "resend":
        _send_resend(sender, to_addr, subject, body, data, fname)
    else:
        from streetlamp.reporting import _send_email_sync

        _send_email_sync(to_addr, subject, body, data, fname)


async def _deliver(sender: dict, to_addr: str, subject: str, body: str, data: bytes, fname: str) -> str | None:
    """성공 시 None, 실패 시 오류 문구."""
    try:
        await asyncio.to_thread(_send_sync, sender, to_addr, subject, body, data, fname)
        return None
    except Exception as e:  # noqa: BLE001 — 수신자별 실패를 모아 안내
        return f"{to_addr}: {type(e).__name__} {str(e)[:200]}"


async def save_sender(session: AsyncSession, form: dict[str, str]) -> None:
    """form: SENDER_DEFAULTS 키. 비밀값은 비어 있으면 기존 값 유지."""
    for key in SENDER_DEFAULTS:
        if key not in form:
            continue
        val = (form[key] or "").strip()
        if key in SECRET_KEYS:
            if val:
                await _set(session, key, seal_secret(val))
            continue
        await _set(session, key, val)


async def send_sender_test(session: AsyncSession, to_addr: str) -> tuple[bool, str]:
    sender = await load_sender(session)
    ok, status = sender_status(sender)
    if not ok:
        return False, status
    now = fmt_kst(datetime.utcnow(), "%Y-%m-%d %H:%M")
    body = (
        "Smart FMS 정비관리 공식 발신 계정 연결 테스트 메일입니다.\n"
        f"발신: {sender_from(sender) or '서버 기본 발신'}\n"
        f"시각(KST): {now}\n\n"
        "이 메일을 받으셨다면 협력사 접수 메일이 같은 계정으로 발송됩니다."
    )
    err = await _deliver(
        sender, to_addr, "[Smart FMS] 발신 계정 연결 테스트", body, build_xlsx([]), "연결테스트.xlsx"
    )
    if err:
        return False, "발송 실패 — " + err
    return True, f"{to_addr} 로 테스트 메일을 보냈습니다. 받은편지함(또는 스팸함)을 확인하세요."


async def run_send(
    session: AsyncSession, *, test: bool = False, now_utc: datetime | None = None
) -> tuple[bool, str]:
    """설정대로 발송 → (실패 없음 여부, 안내 문구). test=True 이면 집계 구간·발송기록을 바꾸지 않음."""
    settings = await load_settings(session)
    sender = await load_sender(session)
    now = now_utc or datetime.utcnow()
    since, until = window_for(settings, now)
    rows = await fetch_receipts(session, since, until)
    stamp = fmt_kst(until, "%Y%m%d_%H%M")
    day = fmt_kst(until, "%Y-%m-%d")
    tag = "[테스트] " if test else ""

    sent: list[str] = []
    errors: list[str] = []
    skipped: list[str] = []

    combined_to, _ = parse_emails(settings["to"])
    if combined_to:
        if rows or settings["send_empty"]:
            data = build_xlsx(rows)
            subject = f"{tag}[정비관리] 협력사 접수 정비의뢰 {len(rows)}건 ({day})"
            body = build_body(rows, since, until)
            for addr in combined_to:
                err = await _deliver(sender, addr, subject, body, data, f"협력사접수_{stamp}.xlsx")
                if err:
                    errors.append(err)
                else:
                    sent.append(addr)
        else:
            skipped.append("통합: 접수 0건")

    if settings["per_partner"]:
        groups: dict[int, list[dict]] = {}
        for r in rows:
            groups.setdefault(r["partner_id"], []).append(r)
        for pid, items in groups.items():
            name = items[0]["partner_name"]
            addrs, _ = parse_emails(settings["partner_to"].get(str(pid)) or items[0]["partner_email"])
            if not addrs:
                skipped.append(f"{name}: 메일 주소 없음")
                continue
            data = build_xlsx(items)
            subject = f"{tag}[정비관리] {name} 접수 정비의뢰 {len(items)}건 ({day})"
            body = build_body(items, since, until, partner_name=name)
            for addr in addrs:
                safe = re.sub(r'[\\/:*?"<>|\s]+', "_", name)
                err = await _deliver(sender, addr, subject, body, data, f"{safe}_접수_{stamp}.xlsx")
                if err:
                    errors.append(err)
                else:
                    sent.append(addr)

    no_target = not combined_to and not settings["per_partner"]
    if no_target:
        msg = "받는 메일 주소가 없습니다. 통합 수신 주소를 입력하거나 협력사별 발송을 켜세요."
    else:
        parts = [f"{tag}집계 {fmt_kst(since)}~{fmt_kst(until)} · 접수 {len(rows)}건"]
        if sent:
            parts.append(f"발송 {len(sent)}통 ({', '.join(sent[:10])}{' 외' if len(sent) > 10 else ''})")
        if skipped:
            parts.append("미발송: " + ", ".join(skipped))
        if errors:
            parts.append("실패: " + " / ".join(errors[:5]))
        if not sent and not errors and not skipped:
            parts.append("발송 대상 없음")
        msg = " · ".join(parts)

    if not test:
        if not errors or sent:
            await _set(session, "last_sent_utc", until.isoformat(timespec="seconds"))
        await _set(session, "last_run_date", day)
    await _set(session, "last_result", msg[:1000])
    await _set(session, "last_result_at", fmt_kst(datetime.utcnow(), "%Y-%m-%d %H:%M:%S"))
    await session.commit()
    print(f"[maint-mail] {msg}", flush=True)
    return (not errors and not no_target), msg


async def job_send(session_factory) -> None:
    async with session_factory() as session:
        settings = await load_settings(session)
        if not settings["enabled"]:
            return
        today = datetime.now(KST).strftime("%Y-%m-%d")
        if settings["last_run_date"] == today:
            return
        await run_send(session)


async def reschedule(scheduler, session_factory) -> None:
    """설정을 읽어 매일 발송 작업을 등록/해제."""
    from apscheduler.triggers.cron import CronTrigger

    async with session_factory() as session:
        settings = await load_settings(session)
    try:
        scheduler.remove_job(JOB_ID)
    except Exception:
        pass
    if not settings["enabled"]:
        return
    scheduler.add_job(
        job_send,
        CronTrigger(hour=settings["hour"], minute=settings["minute"], timezone=KST),
        args=[session_factory],
        id=JOB_ID,
        replace_existing=True,
        misfire_grace_time=3600,
        coalesce=True,
    )


def next_run_label(scheduler) -> str:
    if scheduler is None:
        return ""
    try:
        job = scheduler.get_job(JOB_ID)
    except Exception:
        return ""
    if not job or not job.next_run_time:
        return ""
    return job.next_run_time.astimezone(KST).strftime("%Y-%m-%d %H:%M")
