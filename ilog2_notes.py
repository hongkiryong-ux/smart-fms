# -*- coding: utf-8 -*-
"""점검일지2 특이사항 — 하루 한 행 목록·검색·엑셀보내기 공통."""
from __future__ import annotations

import importlib
from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

WEEKDAYS = ("월", "화", "수", "목", "금", "토", "일")

# (모듈명, 건물 판별 함수, 엑셀 파일명 접두)
NOTES_MODULES: dict[str, tuple[str, str]] = {
    "housing_substation": ("is_housing_substation_building", "주택변전소"),
    "central_control_room": ("is_central_control_room_building", "중앙관제실"),
    "ccr_facility": ("is_ccr_facility_building", "중앙관제실(설비)"),
    "steelworks_hq": ("is_steelworks_hq_building", "제철소본부"),
    "steelworks_hall": ("is_steelworks_hall_building", "제철회관"),
    "human_center": ("is_human_center_building", "휴먼센터"),
    "eoulrim_gym": ("is_eoulrim_gym_building", "어울림체육관"),
    "baegun_dorm": ("is_baegun_dorm_building", "백운생활관"),
    "giga_town": ("is_giga_town_building", "기가타운"),
    "park1538": ("is_park1538_building", "PARK1538"),
    "baegun_art_hall": ("is_baegun_art_hall_building", "백운아트홀"),
    "sub53": ("is_sub53_building", "53서브"),
    "baegundae": ("is_baegundae_building", "백운대"),
    "baegun_shopping": ("is_baegun_shopping_building", "백운쇼핑"),
    "rist": ("is_rist_building", "RIST"),
    "ground_gwangyang": ("is_ground_gwangyang_building", "GROUND광양"),
}

HEADERS = ["No", "날짜", "요일", "시간", "구분", "내용"]


def _weekday(iso: str) -> str:
    try:
        return WEEKDAYS[date.fromisoformat(str(iso)[:10]).weekday()]
    except ValueError:
        return ""


def _as_written(text: str) -> str:
    """일지에 입력한 줄바꿈·들여쓰기는 살리고 앞뒤 빈 줄과 줄 끝 공백만 정리."""
    lines = [line.rstrip() for line in str(text or "").replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    return "\n".join(lines).strip("\n")


def _uniq(values: list[str]) -> list[str]:
    seen: list[str] = []
    for v in values:
        if v and v not in seen:
            seen.append(v)
    return seen


def notes_rows(notes_list: dict[str, Any] | None, q: str = "") -> list[dict[str, Any]]:
    """일자별 특이사항을 하루 한 행으로 묶고(일지에 적힌 줄바꿈 유지) 키워드로 거른다."""
    keyword = (q or "").strip().lower()
    rows: list[dict[str, Any]] = []
    for entry in (notes_list or {}).get("entries") or []:
        iso = str(entry.get("date") or "")
        blocks: list[dict[str, str]] = []
        for item in entry.get("items") or []:
            time = str(item.get("time") or "").strip()
            section = str(item.get("section") or "").strip()
            if section == "특이사항":
                section = ""
            text = _as_written(item.get("text") or "")
            if not text.strip():
                continue
            if keyword and keyword not in f"{time} {section} {text}".lower():
                continue
            blocks.append({"time": time, "section": section, "text": text})
        if not blocks:
            continue
        multi = len(blocks) > 1
        for b in blocks:
            b["label"] = " · ".join(v for v in (b["time"], b["section"]) if v) if multi else ""
        text = "\n".join(
            (f"[{b['label']}]\n{b['text']}" if b["label"] else b["text"]) for b in blocks
        )
        rows.append({
            "no": len(rows) + 1,
            "date": iso,
            "weekday": _weekday(iso),
            "time": "\n".join(_uniq([b["time"] for b in blocks])),
            "section": "\n".join(_uniq([b["section"] for b in blocks])),
            "text": text,
            "blocks": blocks,
        })
    return rows


def _excel_line_count(text: str, width_chars: float) -> int:
    total = 0
    for line in str(text or "").split("\n"):
        w = sum(2 if ord(ch) > 0x2E7F else 1 for ch in line)
        total += max(1, -(-w // max(1, int(width_chars))))
    return max(1, total)


def resolve_module(kind: str, building) -> Any | None:
    spec = NOTES_MODULES.get(kind)
    if not spec:
        return None
    module = importlib.import_module(kind)
    checker = getattr(module, spec[0], None)
    if not callable(checker) or not checker(building):
        return None
    return module


def export_rows_workbook(rows: list[dict[str, Any]], *, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "특이사항"
    ws.cell(1, 1, title).font = Font(bold=True, size=13)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    fill = PatternFill("solid", fgColor="DDEBF7")
    for col, header in enumerate(HEADERS, 1):
        cell = ws.cell(3, col, header)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, row in enumerate(rows, 4):
        values = [row["no"], row["date"], row["weekday"], row["time"], row["section"], row["text"]]
        for col, value in enumerate(values, 1):
            cell = ws.cell(r, col, value)
            cell.border = border
            cell.alignment = Alignment(
                horizontal="left" if col == 6 else "center",
                vertical="top" if col == 6 else "center",
                wrap_text=col >= 4,
            )
        lines = max(
            _excel_line_count(row["text"], 88),
            _excel_line_count(row["time"], 8),
            _excel_line_count(row["section"], 13),
        )
        if lines > 1:
            ws.row_dimensions[r].height = min(409, 15 * lines + 3)
    if not rows:
        ws.cell(4, 6, "(특이사항 없음)")
    for letter, width in zip("ABCDEF", (6, 12, 6, 9, 14, 90)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A4"
    if rows:
        ws.auto_filter.ref = f"A3:F{len(rows) + 3}"
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
