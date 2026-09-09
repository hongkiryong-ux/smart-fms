"""광양 주택지역 공공시설물 현황 — 엑셀 기준 UTF-8 데이터 로더."""
from __future__ import annotations

import io
import json
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "resources" / "gwangyang_public_facilities.json"

GWANGYANG_FACILITIES_PATH = "/admin/dashboard/gwangyang-facilities"
EXCEL_HEADERS = [
    "배치",
    "순번",
    "행종류",
    "구분",
    "건물명",
    "준공일자",
    "취득액(천원)",
    "연면적(㎡)",
    "坪",
    "비고",
]
ROW_KINDS = {"data", "subtotal", "total", "empty"}


def is_gwangyang_ops_site(name: str | None) -> bool:
    n = (name or "").strip()
    return "광양운영" in n


@lru_cache(maxsize=1)
def load_gwangyang_facilities() -> dict[str, Any]:
    if not DATA_PATH.exists():
        return {
            "title": "광양 주택지역 공공시설물 현황",
            "as_of": "",
            "source": "",
            "headers": ["구분", "건물명", "준공일자", "취득액(천원)", "연면적(㎡)", "坪", "비고"],
            "left": [],
            "right": [],
        }
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))


def paired_rows(data: dict[str, Any]) -> list[tuple[dict | None, dict | None]]:
    left = data.get("left") or []
    right = data.get("right") or []
    n = max(len(left), len(right))
    out: list[tuple[dict | None, dict | None]] = []
    for i in range(n):
        L = left[i] if i < len(left) else None
        R = right[i] if i < len(right) else None
        out.append((L, R))
    return out


def export_gwangyang_facilities_xlsx(data: dict[str, Any]) -> bytes:
    """가져오기에 다시 사용할 수 있는 공공시설물 엑셀을 생성한다."""
    wb = Workbook()
    ws = wb.active
    ws.title = "공공시설물 현황"
    ws["A1"], ws["B1"] = "제목", str(data.get("title") or "")
    ws["A2"], ws["B2"] = "기준일", str(data.get("as_of") or "")
    ws["A3"], ws["B3"] = "원본", str(data.get("source") or "")
    for cell in ("B1", "B2", "B3"):
        ws[cell].data_type = "s"
    for cell in ("A1", "A2", "A3"):
        ws[cell].font = Font(bold=True)

    header_fill = PatternFill("solid", fgColor="DBEAFE")
    for col, value in enumerate(EXCEL_HEADERS, 1):
        cell = ws.cell(row=5, column=col, value=value)
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_no = 6
    for side_key, side_label in (("left", "좌"), ("right", "우")):
        for order, item in enumerate(data.get(side_key) or [], 1):
            if item is None:
                values = [side_label, order, "empty", "", "", "", "", "", "", ""]
            else:
                values = [
                    side_label,
                    order,
                    item.get("kind") or "data",
                    item.get("category") or "",
                    item.get("name") or "",
                    item.get("completed") or "",
                    item.get("cost") or "",
                    item.get("area") or "",
                    item.get("pyeong") or "",
                    item.get("note") or "",
                ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(
                    row=row_no,
                    column=col,
                    value=str(value) if value is not None else "",
                )
                cell.data_type = "s"
            row_no += 1

    widths = [8, 8, 12, 16, 28, 15, 18, 16, 12, 30]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=5, column=index).column_letter].width = width
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:J{max(5, row_no - 1)}"
    ws.sheet_view.showGridLines = False

    guide = wb.create_sheet("작성안내")
    guide["A1"] = "공공시설물 현황 가져오기 안내"
    guide["A1"].font = Font(bold=True, size=14)
    guide["A3"] = "배치는 좌/우, 순번은 각 표의 표시 순서입니다."
    guide["A4"] = "행종류는 data, subtotal, total, empty 중 하나를 입력하세요."
    guide["A5"] = "열 제목과 시트 구조를 유지한 뒤 .xlsx 파일로 가져오세요."
    guide.column_dimensions["A"].width = 75

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def import_gwangyang_facilities_xlsx(content: bytes) -> tuple[dict[str, Any], int]:
    """내보내기 양식의 엑셀을 검증해 JSON 데이터로 변환한다."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError("올바른 .xlsx 파일이 아닙니다.") from exc
    if "공공시설물 현황" not in wb.sheetnames:
        raise ValueError("'공공시설물 현황' 시트를 찾을 수 없습니다.")
    ws = wb["공공시설물 현황"]
    headers = [str(ws.cell(5, col).value or "").strip() for col in range(1, 11)]
    if headers != EXCEL_HEADERS:
        raise ValueError("엑셀 열 제목이 내보내기 양식과 다릅니다.")

    grouped: dict[str, dict[int, dict | None]] = {"left": {}, "right": {}}
    imported_count = 0
    for row_no in range(6, ws.max_row + 1):
        values = [ws.cell(row_no, col).value for col in range(1, 11)]
        if not any(value not in (None, "") for value in values):
            continue
        side_raw = str(values[0] or "").strip().lower()
        side = {"좌": "left", "left": "left", "우": "right", "right": "right"}.get(
            side_raw
        )
        if not side:
            raise ValueError(f"{row_no}행의 배치는 좌 또는 우여야 합니다.")
        try:
            order = int(values[1])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{row_no}행의 순번이 올바르지 않습니다.") from exc
        if order < 1 or order in grouped[side]:
            raise ValueError(f"{row_no}행의 순번이 중복되었거나 올바르지 않습니다.")
        kind = str(values[2] or "data").strip().lower()
        if kind not in ROW_KINDS:
            raise ValueError(f"{row_no}행의 행종류가 올바르지 않습니다.")
        if kind == "empty":
            grouped[side][order] = None
            continue
        text = [str(value).strip() if value is not None else "" for value in values]
        grouped[side][order] = {
            "category": text[3],
            "name": text[4],
            "completed": text[5],
            "cost": text[6],
            "area": text[7],
            "pyeong": text[8],
            "note": text[9],
            "kind": kind,
        }
        imported_count += 1

    if imported_count == 0:
        raise ValueError("가져올 시설물 데이터가 없습니다.")

    result = {
        "title": str(ws["B1"].value or "광양 주택지역 공공시설물 현황").strip(),
        "as_of": str(ws["B2"].value or "").strip(),
        "source": str(ws["B3"].value or "").strip(),
        "headers": ["구분", "건물명", "준공일자", "취득액(천원)", "연면적(㎡)", "坪", "비고"],
    }
    for side in ("left", "right"):
        rows = grouped[side]
        max_order = max(rows, default=0)
        result[side] = [rows.get(index) for index in range(1, max_order + 1)]
    return result, imported_count
