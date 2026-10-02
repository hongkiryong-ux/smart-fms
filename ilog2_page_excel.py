# -*- coding: utf-8 -*-
"""점검일지 1일 화면(표 병합·색·입력값)을 그대로 엑셀로 변환."""
from __future__ import annotations

import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MAX_BLOCKS = 300
MAX_ROWS = 800
MAX_COLS = 80
MAX_TEXT = 4000

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEX = re.compile(r"^#?([0-9a-fA-F]{6})$")
_NUMBER = re.compile(r"^-?(0|[1-9]\d*)(\.\d+)?$")


def _text(value: Any) -> str:
    return str(value if value is not None else "")[:MAX_TEXT]


def _cell_value(text: str) -> Any:
    raw = text.strip()
    if _NUMBER.match(raw.replace(",", "")) and len(raw) < 16:
        number = float(raw.replace(",", ""))
        return int(number) if number.is_integer() and "." not in raw else number
    return text


def _fill(color: Any) -> PatternFill | None:
    match = _HEX.match(str(color or "").strip())
    if not match:
        return None
    hex6 = match.group(1).upper()
    if hex6 == "FFFFFF":
        return None
    return PatternFill("solid", start_color=hex6, end_color=hex6)


def _place_table(rows: list[list[dict]]) -> tuple[list[tuple[int, int, dict]], int]:
    """rowspan/colspan 을 반영해 (행, 열, 셀) 위치를 계산한다."""
    occupied: set[tuple[int, int]] = set()
    placed: list[tuple[int, int, dict]] = []
    width = 0
    for r, row in enumerate(rows[:MAX_ROWS]):
        c = 0
        for cell in row[:MAX_COLS]:
            while (r, c) in occupied:
                c += 1
            rs = max(1, min(int(cell.get("rs") or 1), MAX_ROWS))
            cs = max(1, min(int(cell.get("cs") or 1), MAX_COLS))
            for dr in range(rs):
                for dc in range(cs):
                    occupied.add((r + dr, c + dc))
            placed.append((r, c, {**cell, "rs": rs, "cs": cs}))
            c += cs
            width = max(width, c)
    return placed, width


def build_page_workbook(payload: dict) -> bytes:
    blocks = list(payload.get("blocks") or [])[:MAX_BLOCKS]
    tables = [
        _place_table(block.get("rows") or [])
        for block in blocks
        if block.get("type") == "table"
    ]
    span = max([w for _, w in tables] + [8])

    wb = Workbook()
    ws = wb.active
    ws.title = (_text(payload.get("sheet")) or "1일")[:31]
    col_px: dict[int, float] = {}
    row_lines: dict[int, int] = {}
    cur = 1
    table_iter = iter(tables)

    def full_row(text: str, size: int, bold: bool, align: str, fill: str | None = None) -> None:
        nonlocal cur
        ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=span)
        cell = ws.cell(row=cur, column=1, value=text)
        cell.font = Font(size=size, bold=bold)
        cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
        if fill:
            cell.fill = PatternFill("solid", start_color=fill, end_color=fill)
        row_lines[cur] = max(row_lines.get(cur, 1), text.count("\n") + 1)
        ws.row_dimensions[cur].height = 26 if size >= 14 else None
        cur += 1

    title = _text(payload.get("title"))
    if title:
        full_row(title, 16, True, "center")
    subtitle = _text(payload.get("subtitle"))
    if subtitle:
        full_row(subtitle, 10, False, "right")

    prev_kind = ""
    for block in blocks:
        kind = block.get("type")
        if kind == "table" and prev_kind == "table":
            cur += 1
        prev_kind = kind
        if kind == "title":
            full_row(_text(block.get("text")), 14, True, "center")
        elif kind == "section":
            cur += 1 if cur > 1 else 0
            full_row(_text(block.get("text")), 11, True, "left")
        elif kind == "line":
            full_row(_text(block.get("text")), 10, False, "left")
        elif kind == "text":
            label = _text(block.get("label"))
            text = _text(block.get("text"))
            if label:
                full_row(label, 11, True, "left")
            ws.merge_cells(start_row=cur, start_column=1, end_row=cur, end_column=span)
            cell = ws.cell(row=cur, column=1, value=text)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            for c in range(1, span + 1):
                ws.cell(row=cur, column=c).border = _BORDER
            row_lines[cur] = max(3, text.count("\n") + 1)
            cur += 1
        elif kind == "table":
            placed, width = next(table_iter)
            if not placed:
                continue
            height = max(r + cell["rs"] for r, _, cell in placed)
            for r, c, cell in placed:
                top, left = cur + r, 1 + c
                bottom, right = top + cell["rs"] - 1, left + cell["cs"] - 1
                text = _text(cell.get("t"))
                if cell["rs"] > 1 or cell["cs"] > 1:
                    ws.merge_cells(start_row=top, start_column=left, end_row=bottom, end_column=right)
                target = ws.cell(row=top, column=left, value=_cell_value(text) if text else None)
                header = bool(cell.get("h"))
                target.font = Font(size=10, bold=header or bool(cell.get("b")), color=_font_color(cell.get("fg")))
                align = cell.get("al") if cell.get("al") in ("left", "right", "center") else "center"
                target.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
                fill = _fill(cell.get("bg"))
                for rr in range(top, bottom + 1):
                    for cc in range(left, right + 1):
                        part = ws.cell(row=rr, column=cc)
                        part.border = _BORDER
                        if fill:
                            part.fill = fill
                try:
                    px = float(cell.get("w") or 0) / cell["cs"]
                except (TypeError, ValueError):
                    px = 0
                if px > 0:
                    for cc in range(left, right + 1):
                        col_px[cc] = max(col_px.get(cc, 0), px)
                lines = text.count("\n") + 1
                per_row = -(-lines // cell["rs"])
                row_lines[top] = max(row_lines.get(top, 1), per_row)
            cur += height
            del width
        else:
            continue

    for col in range(1, span + 1):
        px = col_px.get(col, 70)
        ws.column_dimensions[get_column_letter(col)].width = max(4.0, min(40.0, px / 7.0))
    for row, lines in row_lines.items():
        if lines > 1 and ws.row_dimensions[row].height is None:
            ws.row_dimensions[row].height = min(400, 14 * lines + 4)

    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape" if span > 12 else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _font_color(color: Any) -> str | None:
    match = _HEX.match(str(color or "").strip())
    if not match:
        return None
    hex6 = match.group(1).upper()
    return None if hex6 in ("000000", "111111", "222222", "333333") else hex6


if __name__ == "__main__":
    sample = {
        "title": "TEST 운영일보",
        "subtitle": "2026-10-03",
        "blocks": [
            {"type": "section", "text": "1. UTILITY"},
            {"type": "table", "rows": [
                [{"t": "전력", "rs": 2, "h": True, "w": 80}, {"t": "중간부하", "cs": 2, "h": True, "w": 140}],
                [{"t": "전일", "h": True, "bg": "#fef9c3"}, {"t": "금일", "h": True}],
                [{"t": "연구동", "h": True}, {"t": "100.5"}, {"t": "101"}],
            ]},
            {"type": "table", "rows": [
                [{"t": "기준", "h": True}, {"t": "380V±10%\n220V±6%", "rs": 2}],
                [{"t": "오전", "h": True}],
            ]},
            {"type": "text", "label": "▶ 특이사항", "text": "점검\n완료"},
        ],
    }
    data = build_page_workbook(sample)
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A1"].value == "TEST 운영일보"
    assert "A5:A6" in {str(m) for m in ws.merged_cells.ranges}
    assert ws["B7"].value == 100.5 and ws["C7"].value == 101
    assert ws["A8"].value is None and ws["B9"].value == "380V±10%\n220V±6%"
    assert "B9:B10" in {str(m) for m in ws.merged_cells.ranges}
    assert ws["A11"].value == "▶ 특이사항" and ws["A12"].value == "점검\n완료"
    print("page excel smoke OK")
