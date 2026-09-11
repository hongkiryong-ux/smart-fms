# ai_analysis.py — Smart FMS 전체 데이터 스냅샷 + 집계/GPT 답변
from __future__ import annotations

import json
import re
from io import BytesIO
from datetime import date, datetime, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from models import (
    Building,
    CentralControlRoomDaily,
    CcrFacilityDaily,
    Consumable,
    D1Plan,
    Equipment,
    Floor,
    HousingSubstationDaily,
    InspectionLogBuilding,
    InspectionLogBuilding2,
    InspectionLogFile,
    MaintenanceRecord,
    MaterialItem,
    Notice,
    Partner,
    PMInspection,
    PMResult,
    PMSchedule,
    ScheduleEvent,
    Site,
    SteelworksHqDaily,
    User,
    WorkOrder,
    WorkOrderStatus,
    Zone,
)

try:
    from streetlamp.models import Lamp, MaintenanceRequest as LampRequest
except Exception:  # pragma: no cover
    Lamp = None
    LampRequest = None

_CONTEXT_MAX_CHARS = 52000


def classify_intent(question: str) -> str:
    q = (question or "").strip().lower()
    rules: list[tuple[str, tuple[str, ...]]] = [
        ("equipment", ("설비", "장비", "자산", "equipment", "건물별 설비")),
        ("work_order", ("정비의뢰", "정비접수", "워크오더", "work order", "cmms", "고장수리")),
        ("pm", ("예방점검", "pm", "점검주기", "점검결과", "지연점검")),
        ("streetlamp", ("가로등", "lamp", "불점등", "시민")),
        ("d1", ("d-1", "d1", "작업허가", "협력사 작업", "시설섹션")),
        ("partner", ("협력사", "업체", "partner")),
        ("inspection_log", ("점검일지", "일지", "엑셀일지")),
        (
            "inspection_log2",
            ("점검일지2", "주택변전소", "중앙관제실", "제철소본부", "운영일보"),
        ),
        ("materials", ("자재", "재고", "소모품", "material")),
        ("notices", ("공지", "notice")),
        ("schedules", ("일정", "캘린더", "schedule")),
        ("overview", ("전체", "현황", "요약", "대시보드", "몇", "얼마", "통계", "총")),
    ]
    scores: dict[str, int] = {}
    for intent, kws in rules:
        score = sum(1 for kw in kws if kw in q)
        if score:
            scores[intent] = score
    if not scores:
        return "overview"
    return max(scores, key=scores.get)


def _today() -> date:
    return date.today()


def _enum_val(v: Any) -> str:
    return v.value if hasattr(v, "value") else str(v or "")


def _clip(text: Any, n: int = 200) -> str:
    s = str(text or "").strip()
    return s if len(s) <= n else s[: n - 1] + "…"


def completed_work_order_excel_requested(question: str) -> bool:
    """정비완료 내역의 엑셀 생성을 요청한 자연어인지 판별."""
    q = (question or "").strip().lower()
    return (
        any(word in q for word in ("정비", "워크오더", "work order", "cmms"))
        and any(
            word in q
            for word in ("완료", "조치", "수리", "complete", "completed", "completion")
        )
        and any(word in q for word in ("엑셀", "excel", "xlsx", "파일"))
    )


def excel_export_requested(question: str) -> bool:
    """GPT 답변·조회 결과를 엑셀로 정리해달라는 요청인지 판별."""
    q = (question or "").strip().lower()
    if not q or completed_work_order_excel_requested(q):
        return False
    if not any(token in q for token in ("엑셀", "excel", "xlsx")):
        return False
    action_tokens = (
        "정리",
        "만들",
        "생성",
        "다운",
        "내려",
        "내보내",
        "추출",
        "뽑아",
        "저장",
        "파일",
        "표로",
        "시트",
        "변환",
        "작성",
        "부탁",
        "해줘",
        "해주세요",
        "해 줘",
        "로 줘",
        "로줘",
    )
    if any(token in q for token in action_tokens):
        return True
    return "엑셀로" in q or "excel로" in q or "to excel" in q


_AI_EXCEL_SETTING_PREFIX = "ai_excel_export."
_EXCEL_MAX_SHEETS = 10
_EXCEL_MAX_COLS = 60
_EXCEL_MAX_ROWS = 5000
_MD_TABLE_BLOCK_RE = re.compile(
    r"(?:^\|[^\n]+\|\s*\n^\|[\s:\-|]+\|\s*\n(?:^\|[^\n]+\|\s*\n?)+)",
    re.MULTILINE,
)


def ai_excel_setting_key(user_id: int) -> str:
    return f"{_AI_EXCEL_SETTING_PREFIX}{int(user_id)}"


def _split_md_row(line: str) -> list[str]:
    return [part.strip() for part in line.strip().strip("|").split("|")]


def extract_markdown_tables(text: str) -> list[dict[str, Any]]:
    """마크다운 표를 엑셀 시트 구조로 추출."""
    sheets: list[dict[str, Any]] = []
    for idx, block in enumerate(_MD_TABLE_BLOCK_RE.findall(text or ""), 1):
        lines = [ln.strip() for ln in block.strip().splitlines() if ln.strip()]
        if len(lines) < 2:
            continue
        headers = _split_md_row(lines[0])
        if not any(headers):
            continue
        rows: list[list[Any]] = []
        for line in lines[2:]:
            cells = _split_md_row(line)
            if len(cells) < len(headers):
                cells.extend([""] * (len(headers) - len(cells)))
            elif len(cells) > len(headers):
                cells = cells[: len(headers)]
            rows.append(cells)
        sheets.append({"title": f"표{idx}", "headers": headers, "rows": rows})
    return sheets


def normalize_excel_export_payload(payload: Any) -> dict[str, Any] | None:
    """GPT/파서 결과를 안전한 엑셀 payload로 정규화."""
    if not isinstance(payload, dict):
        return None
    sheets_in = payload.get("sheets")
    if not isinstance(sheets_in, list):
        return None
    sheets: list[dict[str, Any]] = []
    for raw in sheets_in[:_EXCEL_MAX_SHEETS]:
        if not isinstance(raw, dict):
            continue
        headers_raw = raw.get("headers")
        rows_raw = raw.get("rows")
        if not isinstance(headers_raw, list) or not headers_raw:
            continue
        headers = [str(h if h is not None else "").strip() or f"열{i}" for i, h in enumerate(headers_raw[:_EXCEL_MAX_COLS], 1)]
        if not isinstance(rows_raw, list):
            rows_raw = []
        rows: list[list[Any]] = []
        for row in rows_raw[:_EXCEL_MAX_ROWS]:
            if isinstance(row, dict):
                cells = [row.get(h, "") for h in headers]
            elif isinstance(row, (list, tuple)):
                cells = list(row[: len(headers)])
            else:
                cells = [row]
            if len(cells) < len(headers):
                cells.extend([""] * (len(headers) - len(cells)))
            normalized: list[Any] = []
            for cell in cells:
                if cell is None:
                    normalized.append("")
                elif isinstance(cell, (int, float, bool)):
                    normalized.append(cell)
                else:
                    normalized.append(str(cell))
            rows.append(normalized)
        title = str(raw.get("title") or f"Sheet{len(sheets) + 1}").strip() or f"Sheet{len(sheets) + 1}"
        title = re.sub(r'[\\/*?:\[\]]', "_", title)[:31]
        sheets.append({"title": title, "headers": headers, "rows": rows})
    if not sheets:
        return None
    summary = str(payload.get("summary") or "").strip()
    if not summary:
        total_rows = sum(len(s["rows"]) for s in sheets)
        summary = f"{len(sheets)}개 시트로 총 {total_rows}행을 엑셀로 정리했습니다."
    stem = str(payload.get("filename_stem") or "AI_답변_엑셀").strip() or "AI_답변_엑셀"
    stem = re.sub(r'[\\/:*?"<>|]+', "_", stem)[:80]
    return {"summary": summary, "filename_stem": stem, "sheets": sheets}


def parse_excel_json_response(text: str) -> dict[str, Any] | None:
    """GPT 응답에서 엑셀 JSON을 추출·정규화."""
    raw = (text or "").strip()
    if not raw:
        return None
    fence = re.search(r"```(?:json)?\s*([\s\S]*?)```", raw, re.IGNORECASE)
    if fence:
        raw = fence.group(1).strip()
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError:
        start = raw.find("{")
        end = raw.rfind("}")
        if start < 0 or end <= start:
            return None
        try:
            payload = json.loads(raw[start : end + 1])
        except json.JSONDecodeError:
            return None
    return normalize_excel_export_payload(payload)


def export_ai_sheets_xlsx(sheets: list[dict[str, Any]]) -> bytes:
    """정규화된 sheets 목록을 xlsx 바이트로 생성."""
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="DBEAFE")
    wrap = Alignment(vertical="top", wrap_text=True)
    first = True
    for sheet in sheets:
        title = str(sheet.get("title") or "Sheet1")[:31]
        if first:
            ws = wb.active
            ws.title = title
            first = False
        else:
            ws = wb.create_sheet(title=title)
        headers = list(sheet.get("headers") or [])
        rows = list(sheet.get("rows") or [])
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_no, row in enumerate(rows, 2):
            for col, value in enumerate(row, 1):
                cell = ws.cell(row=row_no, column=col, value=value if value is not None else "")
                cell.alignment = wrap
                if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                    cell.value = "'" + value
        for col, header in enumerate(headers, 1):
            width = max(10, min(42, len(str(header)) + 4))
            if rows:
                sample = max((len(str(r[col - 1])) if col - 1 < len(r) else 0) for r in rows[:50])
                width = max(width, min(42, sample + 2))
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
        if headers:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(headers)).column_letter}{max(1, len(rows) + 1)}"
        ws.sheet_view.showGridLines = False
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


_EXCEL_JSON_RULES = (
    "응답은 반드시 JSON 객체 하나만 출력하세요. 설명 문장·마크다운·코드펜스는 넣지 마세요.\n"
    "형식:\n"
    '{"summary":"한국어 짧은 안내","filename_stem":"파일명짧은제목",'
    '"sheets":[{"title":"시트명","headers":["열1","열2"],"rows":[["값1","값2"]]}]}\n'
    "규칙: 사실에 근거한 표만 넣고, 없으면 sheets는 []로 두며 summary에 이유를 적으세요. "
    "시트명은 31자 이내, 열·행은 과도하게 늘리지 마세요."
)


def call_openai_excel_payload(
    *,
    api_key: str,
    model: str,
    question: str,
    context: dict[str, Any] | None = None,
    source_text: str = "",
) -> dict[str, Any]:
    """질문·이전 답변·FMS 데이터를 엑셀용 JSON payload로 변환."""
    if source_text.strip():
        system = (
            "당신은 Smart FMS AI 답변을 엑셀 표로 구조화하는 도우미입니다. "
            "원본 답변의 사실만 사용하고 추측하지 마세요. "
            + _EXCEL_JSON_RULES
        )
        user_msg = (
            f"사용자 요청:\n{question}\n\n"
            f"원본 답변:\n{source_text.strip()[:20000]}\n\n"
            "위 원본을 엑셀 시트로 정리하세요."
        )
        messages = [
            {"role": "system", "content": system},
            {"role": "user", "content": user_msg},
        ]
    else:
        ctx = context or {}
        system = (
            _build_gpt_system_message(ctx)
            + "\n\n사용자가 엑셀 정리를 요청했습니다. "
            + _EXCEL_JSON_RULES
        )
        user_msg = (
            f"질문:\n{question}\n\n"
            "Smart FMS 데이터만 근거로 엑셀용 JSON을 만드세요."
        )
        messages = [
            {"role": "system", "content": system},
            {"role": "user", "content": user_msg},
        ]
    text = _openai_chat_completion(api_key=api_key, model=model, messages=messages)
    payload = parse_excel_json_response(text)
    if not payload:
        raise RuntimeError("엑셀용 표 구조를 만들지 못했습니다. 질문을 조금 더 구체적으로 해 주세요.")
    return payload


async def save_ai_excel_export(db: AsyncSession, user_id: int, payload: dict[str, Any]) -> None:
    """사용자별 최근 AI 엑셀 payload 저장."""
    from models import AppSetting

    key = ai_excel_setting_key(user_id)
    raw = json.dumps(payload, ensure_ascii=False, default=str)
    row = await db.get(AppSetting, key)
    if row is None:
        db.add(AppSetting(key=key, value=raw))
    else:
        row.value = raw
    await db.commit()


async def load_ai_excel_export(db: AsyncSession, user_id: int) -> dict[str, Any] | None:
    from models import AppSetting

    row = await db.get(AppSetting, ai_excel_setting_key(user_id))
    if not row or not row.value:
        return None
    try:
        data = json.loads(row.value)
    except json.JSONDecodeError:
        return None
    return normalize_excel_export_payload(data)


async def clear_ai_excel_export(db: AsyncSession, user_id: int) -> None:
    from models import AppSetting

    row = await db.get(AppSetting, ai_excel_setting_key(user_id))
    if row is not None:
        await db.delete(row)
        await db.commit()


def _completed_work_order_details_requested(question: str) -> bool:
    q = (question or "").strip().lower()
    return any(word in q for word in ("정비", "워크오더", "work order", "cmms")) and any(
        word in q
        for word in ("완료", "조치", "수리", "complete", "completed", "completion")
    )


async def _count(db: AsyncSession, stmt) -> int:
    return int((await db.execute(stmt)).scalar() or 0)


async def load_completed_work_order_rows(db: AsyncSession) -> list[dict[str, Any]]:
    """최종 완료된 정비의뢰 전체를 엑셀·AI 공통 구조로 조회."""
    rows = (
        await db.execute(
            select(WorkOrder)
            .where(
                WorkOrder.is_active == True,  # noqa: E712
                WorkOrder.status.in_(
                    (
                        WorkOrderStatus.completed,
                        WorkOrderStatus.verified,
                        WorkOrderStatus.closed,
                    )
                ),
                WorkOrder.completion_approval_pending.is_not(True),
            )
            .options(
                selectinload(WorkOrder.equipment)
                .selectinload(Equipment.zone)
                .selectinload(Zone.floor)
                .selectinload(Floor.building)
                .selectinload(Building.site),
                selectinload(WorkOrder.partner),
            )
            .order_by(WorkOrder.completed_at.desc().nullslast(), WorkOrder.id.desc())
        )
    ).scalars().all()
    site_ids = {int(row.site_id) for row in rows if row.site_id}
    site_map: dict[int, str] = {}
    if site_ids:
        site_map = {
            int(site_id): str(name or "")
            for site_id, name in (
                await db.execute(select(Site.id, Site.name).where(Site.id.in_(site_ids)))
            ).all()
        }

    result = []
    for row in rows:
        equipment = row.equipment
        zone = equipment.zone if equipment else None
        floor = zone.floor if zone else None
        building = floor.building if floor else None
        site_name = building.site.name if building and building.site else ""
        if not site_name and row.site_id:
            site_name = site_map.get(int(row.site_id), "")
        result.append(
            {
                "id": row.id,
                "status": _enum_val(row.status),
                "site": site_name,
                "building": building.name if building else "",
                "equipment_code": equipment.code if equipment else "",
                "equipment_name": equipment.name if equipment else "",
                "title": row.title or "",
                "description": row.description or "",
                "cause": row.cause or "",
                "action": row.action or "",
                "parts_used": row.parts_used or "",
                "partner": row.partner.name if row.partner else "",
                "assignee": row.assignee_name or "",
                "requester": row.requester_name or "",
                "priority": row.priority or "",
                "work_type": row.work_type or "",
                "scheduled_date": str(row.scheduled_date or ""),
                "completed_at": str(row.completed_at or ""),
                "work_hours": row.work_hours,
                "cost": row.cost,
                "completion_approved_by": row.completion_approved_by or "",
                "completion_approved_at": str(row.completion_approved_at or ""),
            }
        )
    return result


def export_completed_work_orders_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """정비의뢰·완료내용을 한 행에 정리한 엑셀."""
    wb = Workbook()
    ws = wb.active
    ws.title = "정비완료 내역"
    headers = [
        "번호",
        "상태",
        "사업장",
        "건물",
        "설비코드",
        "설비명",
        "정비의뢰 제목",
        "정비의뢰 내용",
        "고장원인",
        "정비완료 조치내용",
        "사용부품",
        "협력사",
        "담당자",
        "요청자",
        "우선순위",
        "작업유형",
        "예정일",
        "완료일시",
        "작업시간",
        "비용",
        "최종승인자",
        "최종승인일시",
    ]
    fill = PatternFill("solid", fgColor="DBEAFE")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    keys = [
        "id",
        "status",
        "site",
        "building",
        "equipment_code",
        "equipment_name",
        "title",
        "description",
        "cause",
        "action",
        "parts_used",
        "partner",
        "assignee",
        "requester",
        "priority",
        "work_type",
        "scheduled_date",
        "completed_at",
        "work_hours",
        "cost",
        "completion_approved_by",
        "completion_approved_at",
    ]
    wrap = Alignment(vertical="top", wrap_text=True)
    for row_no, item in enumerate(rows, 2):
        for col, key in enumerate(keys, 1):
            value = item.get(key)
            cell = ws.cell(row=row_no, column=col, value=value if value is not None else "")
            cell.alignment = wrap
            if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
                cell.value = "'" + value

    widths = [9, 12, 16, 18, 15, 20, 28, 42, 30, 42, 25, 18, 14, 14, 11, 14, 13, 20, 11, 14, 14, 20]
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:V{max(1, len(rows) + 1)}"
    ws.sheet_view.showGridLines = False
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _summarize_log2_daily(data: dict | None) -> dict[str, Any]:
    data = data or {}
    out: dict[str, Any] = {}
    fac = data.get("facility") or {}
    util = fac.get("utility") or {}
    if util:
        out["utility"] = {
            uid: {
                k: row.get(k, "")
                for k in ("daily", "monthly", "today", "prev")
                if (row or {}).get(k) not in (None, "")
            }
            for uid, row in util.items()
            if isinstance(row, dict)
        }
    elec = data.get("electrical") or {}
    if elec:
        out["electrical_blocks"] = list(elec.keys())
    notes = fac.get("notes") or data.get("notes")
    if notes:
        out["notes"] = _clip(notes, 300)
    return out


async def _latest_daily_rows(db: AsyncSession, model) -> list[dict[str, Any]]:
    rows = (
        await db.execute(
            select(model, Building.name)
            .join(Building, Building.id == model.building_id)
            .order_by(model.building_id, model.log_date.desc())
        )
    ).all()
    seen: set[int] = set()
    out: list[dict[str, Any]] = []
    for row, bname in rows:
        if row.building_id in seen:
            continue
        seen.add(row.building_id)
        out.append(
            {
                "building_id": row.building_id,
                "building": bname,
                "date": row.log_date.isoformat(),
                "summary": _summarize_log2_daily(row.data),
            }
        )
    return out


def _match_buildings_in_question(question: str, buildings: list[dict]) -> list[dict]:
    q = (question or "").strip()
    if not q:
        return []
    matched: list[dict] = []
    for b in buildings:
        name = b.get("name") or ""
        code = b.get("code") or ""
        if name and name in q:
            matched.append(b)
        elif code and code in q:
            matched.append(b)
    hint = _extract_building_hint(question)
    if hint:
        for b in buildings:
            name = b.get("name") or ""
            if hint in name and b not in matched:
                matched.append(b)
    return matched[:8]


def _extract_year_month(question: str) -> tuple[int | None, int | None]:
    q = question or ""
    year = None
    month = None
    ym = re.search(r"(20\d{2})\s*년", q)
    if ym:
        year = int(ym.group(1))
    mm = re.search(r"(\d{1,2})\s*월", q)
    if mm:
        month = int(mm.group(1))
    return year, month


def _compact_housing_monthly_power(report: dict[str, Any]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for sec in report.get("sections") or []:
        breaker_map = {b["meter_id"]: b for b in sec.get("breakers") or []}
        sec_data: dict[str, Any] = {
            "section": sec.get("title", ""),
            "monthly_totals_kwh": [],
            "daily_usage_kwh": [],
        }
        for tot in sec.get("totals") or []:
            mid = tot.get("meter_id")
            br = breaker_map.get(mid) or {}
            usage = tot.get("usage_sum")
            if usage not in (None, ""):
                sec_data["monthly_totals_kwh"].append(
                    {
                        "name": br.get("name", mid),
                        "total": usage,
                        "max_a": tot.get("max_a_max", ""),
                    }
                )
        for day in sec.get("days") or []:
            usages: dict[str, Any] = {}
            for br in day.get("breakers") or []:
                if br.get("usage") not in (None, ""):
                    usages[br.get("name", "")] = br["usage"]
            if usages:
                sec_data["daily_usage_kwh"].append(
                    {"day": day.get("day"), "date": day.get("date"), "usage": usages}
                )
        if sec_data["monthly_totals_kwh"] or sec_data["daily_usage_kwh"]:
            out.append(sec_data)
    return out


async def _gather_housing_monthly_reports(
    db: AsyncSession,
    question: str,
    building_rows: list[Building],
) -> list[dict[str, Any]]:
    """주택변전소 월보(일별·월합계 전력사용량) — 질문 연·월 기준."""
    import calendar

    from housing_substation import fetch_monthly_report_data, is_housing_substation_building

    q = question or ""
    year, month = _extract_year_month(q)
    housing_buildings = [b for b in building_rows if is_housing_substation_building(b)]
    if not housing_buildings:
        return []

    power_kw = any(k in q for k in ("전력", "사용량", "kwh", "kw", "월보", "전기", "전류"))
    housing_kw = any(k in q for k in ("주택변전소", "주택"))
    if not (housing_kw or power_kw or year is not None or month is not None):
        return []

    target_year = year or _today().year
    target_month = month or _today().month
    last_day = calendar.monthrange(target_year, target_month)[1]
    d_from = date(target_year, target_month, 1)
    d_to = date(target_year, target_month, last_day)

    reports_out: list[dict[str, Any]] = []
    for b in housing_buildings:
        row_count = await _count(
            db,
            select(func.count(HousingSubstationDaily.id)).where(
                HousingSubstationDaily.building_id == b.id,
                HousingSubstationDaily.log_date >= d_from,
                HousingSubstationDaily.log_date <= d_to,
            ),
        )
        monthly = await fetch_monthly_report_data(db, b.id, target_year, target_month)
        reports_out.append(
            {
                "building": b.name,
                "building_id": b.id,
                "year": target_year,
                "month": target_month,
                "daily_rows_in_month": row_count,
                "power_usage": _compact_housing_monthly_power(monthly),
            }
        )
    return reports_out


async def gather_context(db: AsyncSession, intent: str, question: str) -> dict[str, Any]:
    """Smart FMS 전체 운영 데이터 스냅샷 (일반질문·GPT 공통)."""
    ctx: dict[str, Any] = {
        "intent": intent,
        "question": (question or "").strip(),
        "as_of": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "data_scope": "full_fms_snapshot",
        "sections": {},
    }
    sec = ctx["sections"]

    site_rows = (
        await db.execute(select(Site).where(Site.is_active == True).order_by(Site.name))  # noqa: E712
    ).scalars().all()
    building_rows = (
        await db.execute(
            select(Building)
            .where(Building.is_active == True)  # noqa: E712
            .options(selectinload(Building.site))
            .order_by(Building.name)
        )
    ).scalars().all()
    building_map = {b.id: b for b in building_rows}

    sec["sites"] = [
        {
            "id": s.id,
            "name": s.name,
            "code": s.code,
            "address": _clip(s.address, 120),
            "manager": s.manager_name or "",
        }
        for s in site_rows
    ]
    sec["buildings"] = [
        {
            "id": b.id,
            "name": b.name,
            "code": b.code,
            "site": b.site.name if b.site else "",
            "manager": b.manager_name or "",
        }
        for b in building_rows
    ]
    sec["question_buildings"] = _match_buildings_in_question(question, sec["buildings"])

    overview = {
        "sites": len(sec["sites"]),
        "buildings": len(sec["buildings"]),
        "equipment_active": await _count(
            db, select(func.count(Equipment.id)).where(Equipment.is_active == True)  # noqa: E712
        ),
        "work_orders_active": await _count(
            db, select(func.count(WorkOrder.id)).where(WorkOrder.is_active == True)  # noqa: E712
        ),
        "pm_schedules": await _count(db, select(func.count(PMSchedule.id))),
        "partners": await _count(db, select(func.count(Partner.id))),
        "inspection_log_files": await _count(db, select(func.count(InspectionLogFile.id))),
        "users_active": await _count(
            db,
            select(func.count(User.id)).where(
                User.is_active == True, User.is_approved == True  # noqa: E712
            ),
        ),
        "material_items": await _count(db, select(func.count(MaterialItem.id))),
        "d1_plans": await _count(db, select(func.count(D1Plan.id))),
        "maintenance_records": await _count(db, select(func.count(MaintenanceRecord.id))),
        "consumables": await _count(db, select(func.count(Consumable.id))),
        "notices_active": await _count(
            db, select(func.count(Notice.id)).where(Notice.is_active == True)  # noqa: E712
        ),
        "schedule_events": await _count(
            db,
            select(func.count(ScheduleEvent.id)).where(ScheduleEvent.is_active == True),  # noqa: E712
        ),
    }
    if Lamp is not None:
        overview["lamps"] = await _count(db, select(func.count(Lamp.id)))
    if LampRequest is not None:
        overview["lamp_requests"] = await _count(db, select(func.count(LampRequest.id)))
    sec["overview"] = overview

    eq_rows = (
        await db.execute(
            select(Building.name, func.count(Equipment.id))
            .select_from(Equipment)
            .join(Zone, Equipment.zone_id == Zone.id)
            .join(Floor, Zone.floor_id == Floor.id)
            .join(Building, Floor.building_id == Building.id)
            .where(Equipment.is_active == True)  # noqa: E712
            .group_by(Building.id, Building.name)
            .order_by(func.count(Equipment.id).desc())
            .limit(30)
        )
    ).all()
    cat_rows = (
        await db.execute(
            select(Equipment.category, func.count(Equipment.id))
            .where(Equipment.is_active == True)  # noqa: E712
            .group_by(Equipment.category)
            .order_by(func.count(Equipment.id).desc())
            .limit(20)
        )
    ).all()
    recent_eq = (
        await db.execute(
            select(Equipment)
            .where(Equipment.is_active == True)  # noqa: E712
            .order_by(Equipment.id.desc())
            .limit(20)
        )
    ).scalars().all()
    sec["equipment"] = {
        "by_building": [{"building": r[0], "count": int(r[1])} for r in eq_rows],
        "by_category": [{"category": r[0] or "기타", "count": int(r[1])} for r in cat_rows],
        "recent": [
            {
                "id": e.id,
                "name": _clip(e.name, 80),
                "code": e.code or "",
                "category": _enum_val(e.category),
                "status": e.status or "",
            }
            for e in recent_eq
        ],
    }

    status_rows = (
        await db.execute(
            select(WorkOrder.status, func.count(WorkOrder.id))
            .where(WorkOrder.is_active == True)  # noqa: E712
            .group_by(WorkOrder.status)
        )
    ).all()
    status_map = {_enum_val(st): int(n) for st, n in status_rows}
    open_n = sum(status_map.get(k, 0) for k in ("received", "assigned", "in_progress"))
    recent_wo = (
        await db.execute(
            select(WorkOrder)
            .where(WorkOrder.is_active == True)  # noqa: E712
            .order_by(WorkOrder.id.desc())
            .limit(25)
        )
    ).scalars().all()
    completion_pending = await _count(
        db,
        select(func.count(WorkOrder.id)).where(
            WorkOrder.is_active == True,  # noqa: E712
            WorkOrder.completion_approval_pending == True,  # noqa: E712
        ),
    )
    sec["work_orders"] = {
        "by_status": status_map,
        "open_count": open_n,
        "completed": status_map.get("completed", 0)
        + status_map.get("verified", 0)
        + status_map.get("closed", 0)
        - completion_pending,
        "completion_approval_pending": completion_pending,
        "approval_pending": await _count(
            db,
            select(func.count(WorkOrder.id)).where(
                WorkOrder.is_active == True,  # noqa: E712
                WorkOrder.approval_requested == True,  # noqa: E712
                WorkOrder.work_permitted == False,  # noqa: E712
            ),
        ),
        "recent": [
            {
                "id": wo.id,
                "title": _clip(wo.title, 80),
                "status": _enum_val(wo.status),
                "priority": wo.priority or "",
                "d1_approved": bool(wo.d1_approved),
                "work_permitted": bool(wo.work_permitted),
                "partner_id": wo.partner_id,
                "assignee": wo.assignee_name or "",
                "scheduled_date": str(wo.scheduled_date or ""),
            }
            for wo in recent_wo
        ],
    }
    if _completed_work_order_details_requested(question):
        completed_rows = await load_completed_work_order_rows(db)
        sec["work_orders"]["completed_details_total"] = len(completed_rows)
        sec["work_orders"]["completed_details"] = [
            {
                **row,
                "description": _clip(row.get("description"), 300),
                "cause": _clip(row.get("cause"), 240),
                "action": _clip(row.get("action"), 400),
                "parts_used": _clip(row.get("parts_used"), 200),
            }
            for row in completed_rows[:50]
        ]

    today = _today()
    due_soon = await _count(
        db,
        select(func.count(PMSchedule.id)).where(
            PMSchedule.is_active == True,  # noqa: E712
            PMSchedule.next_due != None,  # noqa: E711
            PMSchedule.next_due <= today + timedelta(days=7),
        ),
    )
    overdue = await _count(
        db,
        select(func.count(PMSchedule.id)).where(
            PMSchedule.is_active == True,  # noqa: E712
            PMSchedule.next_due != None,  # noqa: E711
            PMSchedule.next_due < today,
        ),
    )
    result_rows = (
        await db.execute(
            select(PMInspection.result, func.count(PMInspection.id)).group_by(PMInspection.result)
        )
    ).all()
    results = {_enum_val(r): int(n) for r, n in result_rows}
    recent_pm = (
        await db.execute(select(PMInspection).order_by(PMInspection.id.desc()).limit(15))
    ).scalars().all()
    sec["pm"] = {
        "schedules": overview["pm_schedules"],
        "due_within_7_days": due_soon,
        "overdue": overdue,
        "inspection_results": results,
        "fault": results.get(PMResult.fault.value, results.get("fault", 0)),
        "caution": results.get(PMResult.caution.value, results.get("caution", 0)),
        "recent_inspections": [
            {
                "id": p.id,
                "result": _enum_val(p.result),
                "inspected_at": str(p.inspected_at or ""),
                "notes": _clip(p.note, 100),
            }
            for p in recent_pm
        ],
    }

    d1_status_rows = (
        await db.execute(select(D1Plan.status, func.count(D1Plan.id)).group_by(D1Plan.status))
    ).all()
    recent_d1 = (
        await db.execute(select(D1Plan).order_by(D1Plan.id.desc()).limit(20))
    ).scalars().all()
    sec["d1_plans"] = {
        "by_status": {_enum_val(st): int(n) for st, n in d1_status_rows},
        "recent": [
            {
                "id": p.id,
                "title": _clip(p.title, 80),
                "work_date": str(p.work_date or ""),
                "status": _enum_val(p.status),
                "partner_id": p.partner_id,
                "building_id": p.building_id,
                "is_urgent": bool(p.is_urgent),
            }
            for p in recent_d1
        ],
    }

    partners = (await db.execute(select(Partner).order_by(Partner.name).limit(50))).scalars().all()
    sec["partners"] = [
        {
            "id": p.id,
            "name": p.name,
            "code": getattr(p, "code", "") or "",
            "contact": _clip(getattr(p, "contact_name", "") or "", 40),
            "risk_grade": getattr(p, "risk_grade", "") or "",
        }
        for p in partners
    ]

    ilog_rows = (
        await db.execute(
            select(InspectionLogBuilding, Building)
            .join(Building, Building.id == InspectionLogBuilding.building_id)
            .order_by(Building.name)
        )
    ).all()
    ilog2_rows = (
        await db.execute(
            select(InspectionLogBuilding2, Building)
            .join(Building, Building.id == InspectionLogBuilding2.building_id)
            .order_by(Building.name)
        )
    ).all()
    ilog_files = (
        await db.execute(
            select(InspectionLogFile.building_id, func.count(InspectionLogFile.id))
            .group_by(InspectionLogFile.building_id)
        )
    ).all()
    file_map = {int(bid): int(n) for bid, n in ilog_files}
    sec["inspection_logs"] = {
        "registered_buildings": [
            {"id": b.id, "name": b.name, "code": b.code, "files": file_map.get(b.id, 0)}
            for _, b in ilog_rows
        ],
        "files_total": overview["inspection_log_files"],
    }
    sec["inspection_logs2"] = {
        "registered_buildings": [
            {"id": b.id, "name": b.name, "code": b.code} for _, b in ilog2_rows
        ],
        "housing_substation": {
            "daily_records": await _count(db, select(func.count(HousingSubstationDaily.id))),
            "latest": await _latest_daily_rows(db, HousingSubstationDaily),
        },
        "central_control_room": {
            "daily_records": await _count(db, select(func.count(CentralControlRoomDaily.id))),
            "latest": await _latest_daily_rows(db, CentralControlRoomDaily),
        },
        "ccr_facility": {
            "daily_records": await _count(db, select(func.count(CcrFacilityDaily.id))),
            "latest": await _latest_daily_rows(db, CcrFacilityDaily),
        },
        "steelworks_hq": {
            "daily_records": await _count(db, select(func.count(SteelworksHqDaily.id))),
            "latest": await _latest_daily_rows(db, SteelworksHqDaily),
        },
    }

    materials = (
        await db.execute(select(MaterialItem).order_by(MaterialItem.name).limit(80))
    ).scalars().all()
    sec["materials"] = {
        "items": [
            {
                "name": m.name,
                "quantity": int(m.quantity or 0),
                "group": m.group_name or "",
                "location": m.location or "",
                "spec": _clip(m.spec, 60),
            }
            for m in materials
        ],
        "low_stock": [
            {"name": m.name, "quantity": int(m.quantity or 0)}
            for m in materials
            if int(m.quantity or 0) <= 5
        ][:20],
    }

    notices = (
        await db.execute(
            select(Notice)
            .where(Notice.is_active == True)  # noqa: E712
            .order_by(Notice.is_pinned.desc(), Notice.published_at.desc())
            .limit(15)
        )
    ).scalars().all()
    sec["notices"] = [
        {
            "id": n.id,
            "title": _clip(n.title, 100),
            "category": n.category or "",
            "pinned": bool(n.is_pinned),
            "published_at": str(n.published_at or ""),
            "body": _clip(n.body, 200),
        }
        for n in notices
    ]

    upcoming = (
        await db.execute(
            select(ScheduleEvent)
            .where(
                ScheduleEvent.is_active == True,  # noqa: E712
                ScheduleEvent.event_date >= today,
                ScheduleEvent.event_date <= today + timedelta(days=60),
            )
            .order_by(ScheduleEvent.event_date.asc())
            .limit(30)
        )
    ).scalars().all()
    sec["schedules"] = [
        {
            "id": e.id,
            "title": _clip(e.title, 80),
            "category": e.category or "",
            "date": str(e.event_date or ""),
            "time": e.event_time or "",
            "location": e.location or "",
        }
        for e in upcoming
    ]

    role_rows = (
        await db.execute(
            select(User.role, func.count(User.id)).where(
                User.is_active == True, User.is_approved == True  # noqa: E712
            ).group_by(User.role)
        )
    ).all()
    sec["users"] = {
        "active_approved": overview["users_active"],
        "by_role": {_enum_val(r): int(n) for r, n in role_rows},
    }

    recent_maint = (
        await db.execute(select(MaintenanceRecord).order_by(MaintenanceRecord.id.desc()).limit(15))
    ).scalars().all()
    sec["maintenance_records"] = {
        "total": overview["maintenance_records"],
        "recent": [
            {
                "id": m.id,
                "equipment_id": m.equipment_id,
                "title": _clip(m.title, 80),
                "work_date": str(m.work_date or ""),
                "action": _clip(m.action, 100),
                "cost": m.cost,
            }
            for m in recent_maint
        ],
    }

    if Lamp is not None:
        lamp_sec: dict[str, Any] = {"lamps": overview.get("lamps", 0)}
        if LampRequest is not None:
            lamp_sec["requests_total"] = overview.get("lamp_requests", 0)
            try:
                req_rows = (
                    await db.execute(
                        select(LampRequest.status, func.count(LampRequest.id)).group_by(
                            LampRequest.status
                        )
                    )
                ).all()
                lamp_sec["requests_by_status"] = {
                    _enum_val(s): int(n) for s, n in req_rows
                }
            except Exception:
                pass
            recent_req = (
                await db.execute(select(LampRequest).order_by(LampRequest.id.desc()).limit(12))
            ).scalars().all()
            lamp_sec["recent_requests"] = [
                {
                    "id": r.id,
                    "status": _enum_val(getattr(r, "status", "")),
                    "note": _clip(
                        (
                            _enum_val(getattr(r, "request_type", ""))
                            + " "
                            + str(getattr(r, "content", "") or "")
                        ),
                        80,
                    ),
                }
                for r in recent_req
            ]
        sec["streetlamp"] = lamp_sec

    housing_monthly = await _gather_housing_monthly_reports(db, question, building_rows)
    if housing_monthly:
        sec["housing_monthly_reports"] = housing_monthly

    return ctx


def _extract_building_hint(question: str) -> str | None:
    q = (question or "").strip()
    m = re.search(r"([가-힣A-Za-z0-9\-]+)\s*(건물|동|센터|관)", q)
    if m:
        return m.group(1)
    return None


def _format_section_lines(title: str, lines: list[str]) -> list[str]:
    if not lines:
        return []
    return ["", f"■ {title}"] + lines


def _answer_for_buildings(sec: dict, matched: list[dict]) -> list[str]:
    lines: list[str] = []
    eq_by_name = {r["building"]: r["count"] for r in sec.get("equipment", {}).get("by_building", [])}
    ilog = {b["name"]: b for b in sec.get("inspection_logs", {}).get("registered_buildings", [])}
    ilog2 = {b["name"]: b for b in sec.get("inspection_logs2", {}).get("registered_buildings", [])}
    for b in matched:
        name = b.get("name", "")
        parts = [f"{name} (코드 {b.get('code', '')}, 사업장 {b.get('site', '')})"]
        if name in eq_by_name:
            parts.append(f"활성 설비 {eq_by_name[name]:,}건")
        if name in ilog:
            parts.append(f"점검일지 등록·파일 {ilog[name].get('files', 0)}건")
        if name in ilog2:
            parts.append("점검일지2 등록됨")
        lines.append("  · " + " · ".join(parts))
    return lines


def build_aggregate_answer(ctx: dict[str, Any], question: str) -> str:
    """전체 FMS 데이터 기반 일반질문 답변."""
    sec = ctx.get("sections", {})
    ov = sec.get("overview", {})
    lines: list[str] = []
    matched = sec.get("question_buildings") or []
    if matched:
        lines.extend(_format_section_lines("질문 관련 건물", _answer_for_buildings(sec, matched)))

    for rep in sec.get("housing_monthly_reports") or []:
        hm_lines = [
            f"  {rep['building']} {rep['year']}년 {rep['month']}월 "
            f"(입력 일지 {rep['daily_rows_in_month']}일)"
        ]
        for sec_pwr in rep.get("power_usage") or []:
            hm_lines.append(f"  [{sec_pwr.get('section', '')}]")
            for t in sec_pwr.get("monthly_totals_kwh") or []:
                hm_lines.append(f"    · {t['name']}: 월합계 {t['total']} kWh")
            for d in sec_pwr.get("daily_usage_kwh") or []:
                parts = ", ".join(f"{k}={v}kWh" for k, v in (d.get("usage") or {}).items())
                hm_lines.append(f"    · {d.get('day')}일: {parts}")
        lines.extend(_format_section_lines("주택변전소 전력 사용량", hm_lines))

    lines.append(f"[Smart FMS 전체 데이터] 기준 시각: {ctx.get('as_of', '')}")
    lines.append(
        f"사업장 {ov.get('sites', 0)} · 건물 {ov.get('buildings', 0)} · "
        f"활성 설비 {ov.get('equipment_active', 0):,} · 정비의뢰 {ov.get('work_orders_active', 0)} · "
        f"PM {ov.get('pm_schedules', 0)} · D-1 {ov.get('d1_plans', 0)} · "
        f"자재 {ov.get('material_items', 0)} · 점검일지 파일 {ov.get('inspection_log_files', 0)}"
        + (f" · 가로등 {ov.get('lamps', 0):,}" if "lamps" in ov else "")
    )

    eq = sec.get("equipment", {})
    eq_lines = [f"  · {r['building']}: {r['count']:,}건" for r in eq.get("by_building", [])[:10]]
    lines.extend(_format_section_lines("설비 (건물별)", eq_lines))

    wo = sec.get("work_orders", {})
    wo_lines = [
        f"  진행 중 {wo.get('open_count', 0)}건 · 완료계열 {wo.get('completed', 0)}건 · "
        f"작업허가 대기 {wo.get('approval_pending', 0)}건"
    ]
    for r in wo.get("recent", [])[:6]:
        wo_lines.append(
            f"  · #{r['id']} [{r['status']}] {r['title']}"
            + (" (D-1승인)" if r.get("d1_approved") else "")
        )
    if "completed_details_total" in wo:
        wo_lines.append(
            f"  정비완료 상세 조회: 총 {wo.get('completed_details_total', 0)}건"
        )
        for row in wo.get("completed_details", [])[:10]:
            wo_lines.append(
                f"  · #{row.get('id')} {row.get('title')}"
                f" / 의뢰: {_clip(row.get('description'), 80) or '-'}"
                f" / 완료: {_clip(row.get('action'), 100) or '-'}"
            )
    lines.extend(_format_section_lines("정비의뢰", wo_lines))

    pm = sec.get("pm", {})
    lines.extend(
        _format_section_lines(
            "예방점검(PM)",
            [
                f"  일정 {pm.get('schedules', 0)} · 7일 이내 {pm.get('due_within_7_days', 0)} · "
                f"지연 {pm.get('overdue', 0)} · 고장 {pm.get('fault', 0)} · 주의 {pm.get('caution', 0)}"
            ],
        )
    )

    d1 = sec.get("d1_plans", {})
    if d1:
        st = d1.get("by_status") or {}
        d1_lines = [f"  상태: {', '.join(f'{k}={v}' for k, v in st.items())}" if st else ""]
        for r in d1.get("recent", [])[:5]:
            d1_lines.append(f"  · {r.get('work_date')} [{r.get('status')}] {r.get('title')}")
        lines.extend(_format_section_lines("D-1 작업계획", [x for x in d1_lines if x]))

    ilog2 = sec.get("inspection_logs2", {})
    if ilog2:
        reg = ", ".join(b["name"] for b in ilog2.get("registered_buildings", [])) or "없음"
        il2_lines = [f"  등록 건물: {reg}"]
        for key, label in (
            ("housing_substation", "주택변전소"),
            ("central_control_room", "중앙관제실(전기)"),
            ("ccr_facility", "중앙관제실(설비)"),
            ("steelworks_hq", "제철소본부"),
        ):
            mod = ilog2.get(key) or {}
            latest = mod.get("latest") or []
            il2_lines.append(f"  · {label}: 일지 {mod.get('daily_records', 0)}건")
            for row in latest[:2]:
                il2_lines.append(f"    - {row.get('building')} 최근 {row.get('date')}")
        lines.extend(_format_section_lines("점검일지2", il2_lines))

    mats = sec.get("materials", {})
    if mats.get("items"):
        mat_lines = [
            f"  · {m['name']}: {m['quantity']}{' (' + m['group'] + ')' if m.get('group') else ''}"
            for m in mats["items"][:12]
        ]
        if mats.get("low_stock"):
            mat_lines.append(
                "  재고 부족(≤5): "
                + ", ".join(f"{m['name']}({m['quantity']})" for m in mats["low_stock"][:8])
            )
        lines.extend(_format_section_lines("자재", mat_lines))

    if sec.get("notices"):
        n_lines = [
            f"  · [{n.get('category')}] {n.get('title')}" for n in sec["notices"][:5]
        ]
        lines.extend(_format_section_lines("공지", n_lines))

    if sec.get("schedules"):
        s_lines = [
            f"  · {e.get('date')} {e.get('title')} ({e.get('category')})"
            for e in sec["schedules"][:8]
        ]
        lines.extend(_format_section_lines("주요설비 일정(60일)", s_lines))

    if sec.get("partners"):
        names = [p["name"] for p in sec["partners"][:15]]
        lines.extend(_format_section_lines("협력사", ["  " + ", ".join(names)]))

    if sec.get("streetlamp"):
        sl = sec["streetlamp"]
        lines.extend(
            _format_section_lines(
                "가로등",
                [f"  등록 {sl.get('lamps', 0):,} · 의뢰 {sl.get('requests_total', 0)}"],
            )
        )

    lines.append("")
    lines.append(
        "※ Smart FMS 전체 DB 스냅샷 기반 답변입니다. 해석·제안이 필요하면 「AI 질문(GPT)」을 사용하세요."
    )
    return "\n".join(lines)


def format_aggregate_answer(ctx: dict[str, Any], *, include_footer: bool = True) -> str:
    """집계 근거 텍스트 (GPT evidence용)."""
    return build_aggregate_answer(ctx, ctx.get("question", "")).replace(
        "\n※ Smart FMS 전체 DB 스냅샷 기반 답변입니다. 해석·제안이 필요하면 「AI 질문(GPT)」을 사용하세요.",
        "\n※ GPT 분석에 사용된 FMS 전체 데이터 요약입니다." if not include_footer else "",
    )


def _sanitize_openai_api_key(api_key: str) -> str:
    key = (api_key or "").strip().replace("\ufeff", "")
    if not key or "…" in key or set(key) <= {"•", "*"}:
        raise RuntimeError(
            "OpenAI API 키가 올바르지 않습니다. "
            "마스킹된 값이 아니라 원본 키(sk-...)를 다시 저장해 주세요."
        )
    try:
        key.encode("ascii")
    except UnicodeEncodeError as e:
        raise RuntimeError(
            "API 키에 한글/유니코드 문자가 포함되어 있습니다. "
            "OpenAI 키(영문·숫자·기호만)를 다시 등록해 주세요."
        ) from e
    return key


def _gpt_system_base() -> str:
    return (
        "당신은 POSCO WIDE Smart FMS 시설관리 분석 도우미입니다. "
        "제공된 JSON은 Smart FMS에 등록된 전체 운영 데이터의 최신 스냅샷입니다 "
        "(사업장·건물·설비·정비의뢰·PM·D-1·협력사·점검일지·점검일지2·자재·공지·일정·가로등 등). "
        "housing_monthly_reports에는 주택변전소 월보 전력사용량(일별·월합계 kWh)이 포함됩니다. "
        "JSON에 있는 수치·목록만 근거로 질문에 한국어로 답하세요. "
        "없는 정보는 추측하지 말고 '데이터에 없음'이라고 하세요. "
        "이전 대화 맥락을 유지하며 후속 질문·추가 설명 요청에도 답하세요. "
        "목록·비교·집계는 가능하면 마크다운 표로 정리하세요. "
        "사용자가 엑셀 정리를 요청하면 시스템이 표 데이터를 파일로 만들어 주므로, "
        "표·목록을 명확히 제시하면 됩니다. "
        "비밀번호·API키·개인 연락처는 언급하지 마세요. "
        "답변 서두에 'GPT 분석'이라고 쓰지 말고 바로 본론부터 작성하세요."
    )


def _build_gpt_system_message(context: dict[str, Any]) -> str:
    payload_ctx = json.dumps(context, ensure_ascii=False, default=str, separators=(",", ":"))
    if len(payload_ctx) > _CONTEXT_MAX_CHARS:
        payload_ctx = payload_ctx[:_CONTEXT_MAX_CHARS] + "..."
    return (
        f"{_gpt_system_base()}\n\n"
        f"데이터 기준 시각: {context.get('as_of', '')}\n"
        f"Smart FMS 전체 데이터(JSON):\n{payload_ctx}"
    )


def _openai_chat_completion(*, api_key: str, model: str, messages: list[dict[str, str]]) -> str:
    import urllib.error
    import urllib.request

    key = _sanitize_openai_api_key(api_key)
    model_name = (model or "gpt-4o-mini").strip() or "gpt-4o-mini"
    try:
        model_name.encode("ascii")
    except UnicodeEncodeError:
        model_name = "gpt-4o-mini"

    body_obj = {
        "model": model_name,
        "temperature": 0.2,
        "messages": messages,
    }

    def _post(use_model: str) -> dict:
        payload = dict(body_obj)
        payload["model"] = use_model
        data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(
            "https://api.openai.com/v1/chat/completions",
            data=data,
            method="POST",
            headers={
                "Authorization": "Bearer " + key,
                "Content-Type": "application/json; charset=utf-8",
                "Accept": "application/json",
                "User-Agent": "SmartFMS/1.0",
            },
        )
        try:
            with urllib.request.urlopen(req, timeout=120) as resp:
                raw = resp.read().decode("utf-8", errors="replace")
                return json.loads(raw)
        except urllib.error.HTTPError as e:
            err_body = e.read().decode("utf-8", errors="replace")[:400]
            if e.code == 404 and use_model != "gpt-4o-mini":
                return _post("gpt-4o-mini")
            raise RuntimeError(f"OpenAI 오류 ({e.code}): {err_body}") from e

    data = _post(model_name)
    text = (
        data.get("choices", [{}])[0]
        .get("message", {})
        .get("content", "")
        .strip()
    )
    if not text:
        raise RuntimeError("OpenAI 응답이 비어 있습니다.")
    return text


def call_openai_conversation(
    *,
    api_key: str,
    model: str,
    context: dict[str, Any],
    chat_messages: list[dict[str, str]],
) -> str:
    """대화형 GPT — 세션 메시지 + FMS 컨텍스트."""
    messages: list[dict[str, str]] = [
        {"role": "system", "content": _build_gpt_system_message(context)}
    ]
    messages.extend(chat_messages[-20:])
    return _openai_chat_completion(api_key=api_key, model=model, messages=messages)


def call_openai_detail(
    *,
    api_key: str,
    model: str,
    question: str,
    context: dict[str, Any],
) -> str:
    user_msg = (
        f"질문:\n{question}\n\n"
        "위 Smart FMS 데이터만 근거로 질문에 답하고, 필요하면 표나 목록으로 정리하세요."
    )
    return call_openai_conversation(
        api_key=api_key,
        model=model,
        context=context,
        chat_messages=[{"role": "user", "content": user_msg}],
    )


async def run_chat_turn(
    db: AsyncSession,
    question: str,
    *,
    chat_messages: list[dict[str, str]] | None,
    api_key: str,
    model: str = "gpt-4o-mini",
) -> dict[str, Any]:
    """GPT 대화 1턴 — 이전 메시지에 이어서 답변."""
    import asyncio

    q = (question or "").strip()
    if not q:
        return {
            "ok": False,
            "needs_api_key": False,
            "messages": list(chat_messages or []),
            "answer": "",
            "evidence": "",
            "intent": "overview",
            "error": "질문을 입력해 주세요.",
        }

    if completed_work_order_excel_requested(q):
        rows = await load_completed_work_order_rows(db)
        history = list(chat_messages or [])
        history.append({"role": "user", "content": q})
        if rows:
            answer = (
                f"최종 완료된 정비의뢰 {len(rows)}건을 정비의뢰 내용과 "
                "정비완료 조치내용이 함께 나오도록 엑셀로 정리했습니다."
            )
            download_url = "/admin/ai-analysis/work-orders/completed.xlsx"
        else:
            answer = "최종 승인된 정비완료 내역이 없어 엑셀을 만들 수 없습니다."
            download_url = ""
        history.append({"role": "assistant", "content": answer})
        return {
            "ok": True,
            "needs_api_key": False,
            "messages": history[-20:],
            "answer": answer,
            "evidence": "",
            "intent": "work_order",
            "error": "",
            "download_url": download_url,
        }

    history = list(chat_messages or [])
    if excel_export_requested(q):
        prior_answer = ""
        for msg in reversed(history):
            if msg.get("role") == "assistant" and str(msg.get("content") or "").strip():
                prior_answer = str(msg.get("content") or "").strip()
                break

        excel_payload = None
        md_sheets = extract_markdown_tables(prior_answer) if prior_answer else []
        if md_sheets:
            excel_payload = normalize_excel_export_payload(
                {
                    "summary": (
                        f"이전 답변의 표 {len(md_sheets)}개를 엑셀로 정리했습니다."
                    ),
                    "filename_stem": "AI_답변_엑셀",
                    "sheets": md_sheets,
                }
            )

        if excel_payload is None:
            key = (api_key or "").strip()
            if not key:
                return {
                    "ok": False,
                    "needs_api_key": True,
                    "messages": history,
                    "answer": "",
                    "evidence": "",
                    "intent": "overview",
                    "error": "OpenAI API 키가 필요합니다.",
                }
            combined_q = " ".join(
                m["content"] for m in history if m.get("role") == "user"
            )
            if combined_q:
                combined_q += " "
            combined_q += q
            intent = classify_intent(combined_q)
            context = (
                {}
                if prior_answer
                else await gather_context(db, intent, combined_q)
            )
            try:
                excel_payload = await asyncio.to_thread(
                    call_openai_excel_payload,
                    api_key=key,
                    model=model or "gpt-4o-mini",
                    question=q,
                    context=context,
                    source_text=prior_answer,
                )
            except Exception as e:
                history.append({"role": "user", "content": q})
                err_answer = f"엑셀 정리에 실패했습니다.\n{e}"
                history.append({"role": "assistant", "content": err_answer})
                return {
                    "ok": False,
                    "needs_api_key": False,
                    "messages": history[-20:],
                    "answer": err_answer,
                    "evidence": "",
                    "intent": intent,
                    "error": str(e),
                }
        else:
            intent = classify_intent(q)

        history.append({"role": "user", "content": q})
        answer = excel_payload["summary"]
        history.append({"role": "assistant", "content": answer})
        return {
            "ok": True,
            "needs_api_key": False,
            "messages": history[-20:],
            "answer": answer,
            "evidence": "",
            "intent": intent,
            "error": "",
            "download_url": "/admin/ai-analysis/export.xlsx",
            "excel_export": excel_payload,
        }

    key = (api_key or "").strip()
    if not key:
        return {
            "ok": False,
            "needs_api_key": True,
            "messages": list(chat_messages or []),
            "answer": "",
            "evidence": "",
            "intent": "overview",
            "error": "OpenAI API 키가 필요합니다.",
        }

    combined_q = " ".join(m["content"] for m in history if m.get("role") == "user")
    if combined_q:
        combined_q += " "
    combined_q += q
    intent = classify_intent(combined_q)
    context = await gather_context(db, intent, combined_q)
    evidence = format_aggregate_answer(context, include_footer=False) if not history else ""

    history.append({"role": "user", "content": q})
    try:
        answer = await asyncio.to_thread(
            call_openai_conversation,
            api_key=key,
            model=model or "gpt-4o-mini",
            context=context,
            chat_messages=history,
        )
    except Exception as e:
        history.pop()
        return {
            "ok": False,
            "needs_api_key": False,
            "messages": history,
            "answer": f"GPT 호출에 실패했습니다.\n{e}",
            "evidence": evidence,
            "intent": intent,
            "error": str(e),
        }

    history.append({"role": "assistant", "content": answer})
    if len(history) > 20:
        history = history[-20:]

    return {
        "ok": True,
        "needs_api_key": False,
        "messages": history,
        "answer": answer,
        "evidence": evidence,
        "intent": intent,
        "error": "",
    }


async def run_analysis(
    db: AsyncSession,
    question: str,
    *,
    mode: str,
    api_key: str = "",
    model: str = "gpt-4o-mini",
) -> dict[str, Any]:
    import asyncio

    q = (question or "").strip()
    if not q:
        return {
            "ok": False,
            "mode": mode,
            "needs_api_key": False,
            "answer": "질문을 입력해 주세요.",
            "evidence": "",
            "intent": "overview",
            "context": {},
            "error": "질문을 입력해 주세요.",
        }

    intent = classify_intent(q)
    context = await gather_context(db, intent, q)
    evidence = format_aggregate_answer(context, include_footer=False)
    aggregate_text = build_aggregate_answer(context, q)

    if mode != "detail":
        return {
            "ok": True,
            "mode": "aggregate",
            "needs_api_key": False,
            "answer": aggregate_text,
            "evidence": "",
            "intent": intent,
            "context": context,
            "error": "",
        }

    key = (api_key or "").strip()
    if not key:
        return {
            "ok": True,
            "mode": "needs_key",
            "needs_api_key": True,
            "answer": (
                "OpenAI API 키가 없어 GPT 분석을 실행하지 못했습니다.\n"
                "아래 「API 키 등록」에서 키를 저장한 뒤 AI 질문을 다시 눌러 주세요."
            ),
            "evidence": evidence,
            "intent": intent,
            "context": context,
            "error": "",
        }

    try:
        detail = await asyncio.to_thread(
            call_openai_detail,
            api_key=key,
            model=model or "gpt-4o-mini",
            question=q,
            context=context,
        )
        return {
            "ok": True,
            "mode": "detail",
            "needs_api_key": False,
            "answer": detail,
            "evidence": evidence,
            "intent": intent,
            "context": context,
            "error": "",
        }
    except Exception as e:
        return {
            "ok": False,
            "mode": "detail_error",
            "needs_api_key": False,
            "answer": f"GPT 호출에 실패했습니다.\n{e}",
            "evidence": evidence,
            "intent": intent,
            "context": context,
            "error": str(e),
        }
