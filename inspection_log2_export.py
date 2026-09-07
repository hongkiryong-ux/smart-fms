# -*- coding: utf-8 -*-
"""점검일지(운영일보) 공통 엑셀보내기 — 유틸리티 검침 표 기반."""
from __future__ import annotations

from datetime import date
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def meter_label(meter: dict[str, Any]) -> str:
    for key in ("label", "item", "name"):
        val = str(meter.get(key) or "").strip()
        if val:
            return val
    cat = str(meter.get("cat") or "").strip()
    building = str(meter.get("building") or "").strip()
    joined = " / ".join(p for p in (cat, building) if p)
    if joined:
        return joined
    return str(meter.get("id") or "")


def write_sheet_table(
    ws,
    *,
    title: str | None,
    headers: list[str],
    rows: list[list[Any]],
    start_row: int = 1,
) -> None:
    row = start_row
    if title:
        ws.cell(row, 1, title).font = Font(bold=True, size=13)
        row += 1
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row, col, header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    row += 1
    for values in rows:
        for col, value in enumerate(values, 1):
            ws.cell(row, col, "" if value is None else value)
        row += 1
    for col in range(1, max(len(headers), 1) + 1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = 18 if col <= 2 else 12


def export_daily_utility_workbook(
    *,
    title: str,
    log_date: date,
    meter_defs: list[dict[str, Any]],
    utility: dict[str, Any],
    notes: str = "",
    multipliers: dict[str, Any] | None = None,
    extra_sheets: list[tuple[str, list[str], list[list[Any]]]] | None = None,
    extra_rows: list[list[Any]] | None = None,
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "1일"
    headers = ["구분", "항목", "전일지침", "금일지침", "일사용", "월누계", "비고"]
    rows: list[list[Any]] = []
    for meter in meter_defs:
        uid = str(meter.get("id") or "")
        cell = utility.get(uid) or {}
        rows.append([
            str(meter.get("cat") or ""),
            meter_label(meter),
            cell.get("prev", ""),
            cell.get("today", ""),
            cell.get("daily", ""),
            cell.get("monthly", ""),
            cell.get("remark", ""),
        ])
    if extra_rows:
        rows.extend(extra_rows)
    write_sheet_table(
        ws,
        title=f"{title} · {log_date.isoformat()}",
        headers=headers,
        rows=rows,
    )
    if multipliers:
        ws2 = wb.create_sheet("배율")
        mrows = [[str(k), str(v)] for k, v in multipliers.items()]
        write_sheet_table(ws2, title="배율", headers=["키", "값"], rows=mrows)
    if notes:
        ws3 = wb.create_sheet("특이사항")
        write_sheet_table(ws3, title="특이사항", headers=["내용"], rows=[[notes]])
    for sheet_name, headers2, rows2 in extra_sheets or []:
        ws_x = wb.create_sheet(sheet_name[:31] or "추가")
        write_sheet_table(ws_x, title=sheet_name, headers=headers2, rows=rows2)
    return workbook_to_bytes(wb)


def export_monthly_utility_workbook(monthly_report: dict[str, Any], *, title: str) -> bytes:
    meters = list(monthly_report.get("meters") or [])
    days = list(monthly_report.get("days") or [])
    totals = monthly_report.get("totals") or {}
    year = monthly_report.get("year")
    month = monthly_report.get("month")
    wb = Workbook()
    ws = wb.active
    ws.title = "월보"
    headers = ["일"] + [meter_label(m) for m in meters]
    if monthly_report.get("power_sum_total") is not None or any(
        "power_sum" in (d or {}) for d in days
    ):
        headers.append("전력합계")
        include_power = True
    else:
        include_power = False
    rows: list[list[Any]] = []
    for day in days:
        util = day.get("utility") or {}
        row = [day.get("day", "")]
        for meter in meters:
            uid = str(meter.get("id") or "")
            cell = util.get(uid) or {}
            row.append(cell.get("daily", "") if isinstance(cell, dict) else cell)
        if include_power:
            row.append(day.get("power_sum", ""))
        rows.append(row)
    total_row: list[Any] = ["계"]
    for meter in meters:
        uid = str(meter.get("id") or "")
        total_row.append(totals.get(uid, ""))
    if include_power:
        power_total = monthly_report.get("power_sum_total")
        if power_total in (None, ""):
            acc = 0.0
            any_val = False
            for day in days:
                raw = str(day.get("power_sum") or "").replace(",", "")
                try:
                    acc += float(raw)
                    any_val = True
                except ValueError:
                    pass
            power_total = f"{acc:g}" if any_val else ""
        total_row.append(power_total)
    rows.append(total_row)
    write_sheet_table(
        ws,
        title=f"{title} · 월보 {year}-{int(month or 0):02d}",
        headers=headers,
        rows=rows,
    )
    return workbook_to_bytes(wb)


def export_yearly_utility_workbook(yearly_report: dict[str, Any], *, title: str) -> bytes:
    meters = list(yearly_report.get("meters") or [])
    months = list(yearly_report.get("months") or [])
    totals = yearly_report.get("totals") or {}
    year = yearly_report.get("year")
    wb = Workbook()
    ws = wb.active
    ws.title = "년보"
    headers = ["월"] + [meter_label(m) for m in meters]
    rows: list[list[Any]] = []
    for month in months:
        util = month.get("utility") or {}
        row = [month.get("month", "")]
        for meter in meters:
            uid = str(meter.get("id") or "")
            row.append(util.get(uid, ""))
        rows.append(row)
    total_row: list[Any] = ["계"]
    for meter in meters:
        uid = str(meter.get("id") or "")
        total_row.append(totals.get(uid, ""))
    rows.append(total_row)
    write_sheet_table(
        ws,
        title=f"{title} · 년보 {year}",
        headers=headers,
        rows=rows,
    )
    return workbook_to_bytes(wb)


def export_notes_workbook(notes_report: dict[str, Any], *, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "특이사항"
    year = notes_report.get("year")
    month = notes_report.get("month")
    rows: list[list[Any]] = []
    for entry in notes_report.get("entries") or []:
        for item in entry.get("items") or []:
            rows.append([
                entry.get("date", ""),
                item.get("time", ""),
                item.get("section", ""),
                item.get("text", ""),
            ])
    write_sheet_table(
        ws,
        title=f"{title} · 특이사항 {year}-{int(month or 0):02d}",
        headers=["날짜", "시간", "구분", "내용"],
        rows=rows or [["", "", "", "(없음)"]],
    )
    return workbook_to_bytes(wb)


def build_extra_sheets(schema: dict[str, Any], data: dict[str, Any]) -> list[tuple[str, list[str], list[list[Any]]]]:
    sheets: list[tuple[str, list[str], list[list[Any]]]] = []
    run = data.get("equipment_run")
    if isinstance(run, dict) and run:
        defs: list[dict[str, Any]] = []
        section = schema.get("equipment_run") or {}
        for block in section.get("blocks") or []:
            for row in block.get("rows") or []:
                defs.append({**row, "cat": block.get("cat") or block.get("label") or block.get("group") or ""})
        for row in section.get("rows") or []:
            defs.append({**row, "cat": row.get("group") or row.get("cat") or ""})
        if not defs:
            defs = [{"id": k, "label": k} for k in run.keys()]
        rows: list[list[Any]] = []
        for item in defs:
            uid = str(item.get("id") or "")
            cell = run.get(uid) or {}
            rows.append([
                str(item.get("cat") or ""),
                str(item.get("label") or item.get("name") or uid),
                cell.get("run1", ""),
                cell.get("run2", ""),
                cell.get("run3", ""),
                cell.get("today_hrs", ""),
                cell.get("prev_cum", ""),
                cell.get("monthly_cum", ""),
                cell.get("remark", ""),
            ])
        sheets.append((
            "가동시간",
            ["구분", "항목", "1회", "2회", "3회", "금일", "전일누계", "월누계", "비고"],
            rows,
        ))
    heating = data.get("heating_run")
    if isinstance(heating, dict) and heating:
        rows = []
        for uid, cell in heating.items():
            if not isinstance(cell, dict):
                continue
            rows.append([
                uid,
                cell.get("run1", cell.get("today", "")),
                cell.get("prev_cum", cell.get("prev", "")),
                cell.get("monthly_cum", cell.get("monthly", "")),
                cell.get("remark", ""),
            ])
        if rows:
            sheets.append((
                "난방가동",
                ["항목", "금일", "전일누계", "월누계", "비고"],
                rows,
            ))
    return sheets


def build_extra_rows(data: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    if data.get("power_sum_daily") not in (None, ""):
        rows.append([
            "합계",
            "전력합계",
            "",
            "",
            data.get("power_sum_daily", ""),
            data.get("power_sum_monthly", ""),
            "",
        ])
    return rows
